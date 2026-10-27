"""word_report with the firm's own report as the template: the document is
kept as a form, header fields filled, old items and prose gone, this visit's
ITEM blocks cloned from the template's paragraphs."""

from datetime import datetime
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from sitescribe import word_report
from sitescribe.models import Observation, PlanPin, PlanSheet, Project, Severity

EXAMPLES = Path(__file__).resolve().parent.parent / "example data"


def _img(path: Path, size=(400, 300), color="gray") -> Path:
    Image.new("RGB", size, color).save(path)
    return path


def _firm_report(path: Path) -> Path:
    """A past report shaped like a GFDS Structural Observation Record."""
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "99 Green St"
    doc.sections[0].header.add_paragraph("April 17, 2026")
    doc.sections[0].footer.paragraphs[0].text = "ACME Engineers"
    doc.add_paragraph("Structural Observation Record")
    t = doc.add_table(rows=2, cols=4)
    for cell, text in zip(t.rows[0].cells, ("Date:", "04/17/26", "Job No:", "260XX.00")):
        cell.text = text
    for cell, text in zip(t.rows[1].cells, ("Project:", "Old Job", "By:", "Pengyu Chen")):
        cell.text = text
    t.rows[1].cells[0].add_paragraph("Site Visit No:     2")
    doc.add_paragraph("Stage of Construction:").runs[0].bold = True
    p = doc.add_paragraph("Framing above the first floor was mostly complete.")
    p.paragraph_format.left_indent = Inches(0.5)
    doc.add_paragraph("Structural Elements Observed:").runs[0].bold = True
    p = doc.add_paragraph("The structural framing was observed to be in general conformance.")
    p.paragraph_format.left_indent = Inches(0.5)
    doc.add_paragraph("Specific Items Observed").runs[0].bold = True
    for n, text, action in ((1, "Sill screws were not installed.", "The Contractor to install sill screws."),
                            (2, "Old wall evaluated.", "The Contractor will review the detail.")):
        p = doc.add_paragraph()
        p.add_run(f"ITEM 2.{n}").bold = True
        p.add_run("\t" + text)
        p.paragraph_format.left_indent = Inches(1.5)
        p.paragraph_format.first_line_indent = Inches(-1.0)
        doc.add_paragraph()
        pic = doc.add_paragraph()
        pic.paragraph_format.left_indent = Inches(1.5)
        pic.add_run().add_picture(str(_img(path.parent / f"old{n}.jpg")), width=Inches(2))
        a = doc.add_paragraph("Action:\t" + action)
        a.paragraph_format.left_indent = Inches(1.5)
        a.paragraph_format.first_line_indent = Inches(-1.0)
        doc.add_paragraph()
    doc.add_paragraph("Previous Action Items:").runs[0].bold = True
    p = doc.add_paragraph("ITEM 1.1\t")
    p.paragraph_format.left_indent = Inches(1.5)
    doc.add_paragraph("Status: Closed.")
    doc.add_paragraph("Action items in this report require the Contractor to complete the work.")
    doc.add_paragraph("End of Report")
    doc.save(str(path))
    return path


def _project(tmp_path: Path, **kw) -> Project:
    sheet = PlanSheet("S-201 Second Floor Framing Plan", tmp_path / "plans.pdf", 0,
                      _img(tmp_path / "s201.png", (1700, 1100), "white"), 1700, 1100)
    a = _img(tmp_path / "003_a_1.jpg", (300, 400))
    b = _img(tmp_path / "003_b_1.jpg", (400, 300))
    c = _img(tmp_path / "003_b_2.jpg", (300, 400))
    obs = [
        Observation(photo_path=a, photos=[a], description="Discontinued studs at window",
                    severity=Severity.MAJOR, tags=["Framing"],
                    action_item="Provide full height studs per 4/S8.1.",
                    pin=PlanPin(sheet.id, 0.25, 0.40), number=3,
                    taken_at=datetime(2026, 10, 27, 9, 5)),
        Observation(photo_path=b, photos=[b, c], description="Holdown rod in (E) 4X beam",
                    severity=Severity.SAFETY,
                    action_item="Per field coordination:\n1) Install stud both sides.\n2) Provide holdown below.",
                    captions={"003_b_1.jpg": "Looking north"},
                    pin=PlanPin(sheet.id, 0.6, 0.7), number=3),
        Observation(description="Note without a photo", number=5),
        Observation(description="Captured but left out", number=6, include=False),
    ]
    fields = dict(name="Hackathon Site Visit", visit_date=datetime(2026, 10, 27), plans=[sheet],
                  observations=obs, summary="Framing about 80% complete.",
                  template_path=_firm_report(tmp_path / "firm_report.docx"))
    fields.update(kw)
    return Project(**fields)


def _text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def _pictures(doc) -> int:
    return len(list(doc.element.body.iter(qn("wp:inline"))))


def _cells(doc, row: int) -> list[str]:
    return [c.text for c in doc.tables[0].rows[row].cells]


def test_template_kind(tmp_path):
    assert word_report.template_kind(_firm_report(tmp_path / "f.docx")) == "report"
    plain = Document()
    plain.add_paragraph("Letterhead only")
    plain.save(str(tmp_path / "plain.docx"))
    assert word_report.template_kind(tmp_path / "plain.docx") == "styles"
    tagged = Document()
    tagged.add_paragraph("{{ project_name }}")
    tagged.save(str(tmp_path / "tagged.docx"))
    assert word_report.template_kind(tmp_path / "tagged.docx") == "tagged"


def test_form_is_kept_and_old_report_text_is_gone(tmp_path):
    doc = Document(str(word_report.build_report(_project(tmp_path), tmp_path / "r.docx")))
    text = _text(doc)

    # the form: title, section labels and the firm's closing paragraphs stay
    for kept in ("Structural Observation Record", "Stage of Construction:",
                 "Structural Elements Observed:", "Specific Items Observed",
                 "Action items in this report require the Contractor", "End of Report"):
        assert kept in text, kept
    # nothing the old report said survives
    for gone in ("Sill screws", "Old wall evaluated", "install sill screws", "will review the detail",
                 "Framing above the first floor", "general conformance", "Previous Action Items",
                 "Status: Closed", "ITEM 2.1", "ITEM 1.1", "Old Job", "Pengyu Chen", "260XX.00"):
        assert gone not in text, gone
    # the old photos are gone too: 3 photos + 1 plan sheet
    assert _pictures(doc) == 4


def test_header_fields_are_filled_or_blank_never_inherited(tmp_path):
    project = _project(tmp_path, engineer="Bailey Bond", report_options={"job_no": "26101.00"})
    doc = Document(str(word_report.build_report(project, tmp_path / "r.docx")))
    assert _cells(doc, 0) == ["Date:", "10/27/26", "Job No:", "26101.00"]
    assert _cells(doc, 1)[1:] == ["Hackathon Site Visit", "By:", "Bailey Bond"]
    assert "Site Visit No:" in _cells(doc, 1)[0] and "2" not in _cells(doc, 1)[0]

    # the app knows no job number or engineer: blank, not the old report's
    doc = Document(str(word_report.build_report(_project(tmp_path), tmp_path / "r2.docx")))
    assert _cells(doc, 0) == ["Date:", "10/27/26", "Job No:", ""]
    assert _cells(doc, 1)[1:] == ["Hackathon Site Visit", "By:", ""]


def test_running_header_gets_this_visit(tmp_path):
    doc = Document(str(word_report.build_report(_project(tmp_path), tmp_path / "r.docx")))
    header = [p.text for p in doc.sections[0].header.paragraphs]
    assert header == ["Hackathon Site Visit", "October 27, 2026"]
    # a footer with no date is the firm's, untouched
    assert doc.sections[0].footer.paragraphs[0].text == "ACME Engineers"


def test_items_follow_the_template_layout(tmp_path):
    doc = Document(str(word_report.build_report(_project(tmp_path), tmp_path / "r.docx")))
    paras = [p for p in doc.paragraphs if p.text.strip()]
    texts = [p.text for p in paras]
    i = texts.index("ITEM 3.1\tDiscontinued studs at window")
    assert texts[i + 1] == "Severity:\tMajor      Tags: Framing"
    assert texts[i + 2] == "Location:\tS-201 Second Floor Framing Plan"
    assert texts[i + 3] == "Action:\tProvide full height studs per 4/S8.1."
    assert "ITEM 3.2\tHoldown rod in (E) 4X beam" in texts
    assert "Looking north\t" in texts
    assert "1)\tInstall stud both sides." in texts and "2)\tProvide holdown below." in texts
    assert "ITEM 5\tNote without a photo" in texts
    assert "Captured but left out" not in texts and "ITEM 6" not in texts
    assert "Plan Sheets" in texts and "S-201 Second Floor Framing Plan" in texts
    # cloned from the template's ITEM paragraph: same hanging indent
    item = paras[i]
    assert item.paragraph_format.left_indent == Inches(1.5)
    assert item.paragraph_format.first_line_indent == Inches(-1.0)
    assert item.runs[0].bold and item.runs[0].text == "ITEM 3.1"
    # the closing text comes after everything we wrote
    assert texts.index("End of Report") > texts.index("Plan Sheets")


def test_stage_comes_from_the_summary_and_is_blank_without_one(tmp_path):
    doc = Document(str(word_report.build_report(_project(tmp_path), tmp_path / "r.docx")))
    texts = [p.text for p in doc.paragraphs]
    assert texts[texts.index("Stage of Construction:") + 1] == "Framing about 80% complete."

    doc = Document(str(word_report.build_report(_project(tmp_path, summary=""), tmp_path / "r2.docx")))
    texts = [p.text.strip() for p in doc.paragraphs]
    i = texts.index("Stage of Construction:")
    assert texts[i + 1] == "" and texts[i + 2] == "Structural Elements Observed:"


def test_report_options_fill_the_prose_and_previous_items(tmp_path):
    project = _project(tmp_path, report_options={
        "stage_of_construction": "Roof framing in progress.",
        "structural_elements_observed": "Roof framing and shear walls.",
        "previous_items": ["ITEM 2.1: Closed", {"number": "2.2", "status": "Open", "description": "Wall"}],
    })
    doc = Document(str(word_report.build_report(project, tmp_path / "r.docx")))
    text = _text(doc)
    assert "Roof framing in progress." in text and "Framing about 80%" not in text
    assert "Roof framing and shear walls." in text
    assert "Previous Action Items:" in text
    assert "ITEM 2.1" in text and "Closed" in text and "ITEM 2.2\tWall" in text and "Open" in text


def test_no_plan_sheets_when_turned_off(tmp_path):
    project = _project(tmp_path, report_options={"include_marked_plans": False})
    doc = Document(str(word_report.build_report(project, tmp_path / "r.docx")))
    assert "Plan Sheets" not in _text(doc) and _pictures(doc) == 3


def test_same_input_gives_same_text(tmp_path):
    project = _project(tmp_path)
    a = _text(Document(str(word_report.build_report(project, tmp_path / "a.docx"))))
    b = _text(Document(str(word_report.build_report(project, tmp_path / "b.docx"))))
    assert a == b


@pytest.mark.parametrize("name", ["260226.Struc Observation.docx", "260417.Struc Observation.docx"])
def test_real_gfds_report_as_template(tmp_path, name):
    template = EXAMPLES / name
    if not template.is_file():
        pytest.skip("GFDS example reports are not in the checkout")
    project = _project(tmp_path, template_path=template)
    doc = Document(str(word_report.build_report(project, tmp_path / "r.docx")))
    text = _text(doc)
    assert word_report.template_kind(template) == "report"
    assert "ITEM 3.1\tDiscontinued studs at window" in text
    assert "End of Report" in text
    for gone in ("Pengyu Chen", "260XX.00", "Sill screws", "Discontinued wall studs", "Status: Closed",
                 "general conformance", "SDS screws", "99 Green St"):
        assert gone not in text, gone
    assert _cells(doc, 0)[:2] == ["Date:", "10/27/26"]
    header = " ".join(p.text for s in doc.sections for p in s.header.paragraphs)
    assert "Hackathon Site Visit" in header and "October 27, 2026" in header and "April" not in header
    assert _pictures(doc) == 4
