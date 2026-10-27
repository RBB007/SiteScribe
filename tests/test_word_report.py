from datetime import datetime
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from sitescribe import word_report
from sitescribe.models import Observation, PlanPin, PlanSheet, Project, Severity


def _img(path: Path, size=(400, 300), color="gray") -> Path:
    Image.new("RGB", size, color).save(path)
    return path


def _project(tmp_path: Path, **kw) -> Project:
    sheet = PlanSheet("S-201 Second Floor Framing Plan", tmp_path / "plans.pdf", 0,
                      _img(tmp_path / "s201.png", (1700, 1100), "white"), 1700, 1100)
    a = _img(tmp_path / "003_a_1.jpg", (300, 400))
    b = _img(tmp_path / "003_b_1.jpg", (400, 300))
    c = _img(tmp_path / "003_b_2.jpg", (300, 400))
    d = _img(tmp_path / "005_a_1.jpg")
    obs = [
        Observation(photo_path=a, photos=[a], description="Discontinued studs at window",
                    severity=Severity.MAJOR, tags=["Framing"],
                    action_item="Contractor to provide full height studs per 4/S8.1.",
                    pin=PlanPin(sheet.id, 0.25, 0.40), number=3,
                    taken_at=datetime(2026, 10, 27, 9, 5)),
        Observation(photo_path=b, photos=[b, c], description="Holdown rod in (E) 4X beam",
                    severity=Severity.SAFETY,
                    action_item="Per field coordination, the Contractor will:\n"
                                "1) Install stud on both sides of the rod.\n"
                                "2) Provide holdown below the (E) 4X beam.",
                    captions={"003_b_1.jpg": "Looking north"},
                    pin=PlanPin(sheet.id, 0.6, 0.7), number=3,
                    taken_at=datetime(2026, 10, 27, 9, 20)),
        Observation(photo_path=d, photos=[d], description="Sheathing nailing OK",
                    severity=Severity.INFO, number=5, taken_at=datetime(2026, 10, 27, 10, 0)),
        Observation(description="Captured but left out", number=6, include=False),
    ]
    fields = dict(name="99 Green St Remodel", site_address="99 Green St, San Francisco",
                  engineer="Pengyu Chen", visit_date=datetime(2026, 10, 27), plans=[sheet],
                  observations=obs, summary="Framing about 80% complete.\n\nShear walls sheathed.",
                  report_options={"job_no": "260XX.00"})
    fields.update(kw)
    return Project(**fields)


def _text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def _picture_count(doc) -> int:
    return len(list(doc.element.body.iter(qn("wp:inline"))))


def test_plain_report_has_all_sections(tmp_path):
    out = tmp_path / "out" / "report.docx"
    result = word_report.build_report(_project(tmp_path), out)

    assert result == out and out.is_file()
    doc = Document(str(out))
    text = _text(doc)
    for expected in ("Site Visit Report", "99 Green St Remodel", "260XX.00", "Summary",
                     "Framing about 80% complete.", "Plan Sheets", "Observations",
                     "ITEM 3.1", "ITEM 3.2", "ITEM 5", "Location:\tS-201 Second Floor Framing Plan",
                     "Looking north", "Action Items", "Photo Log", "1) Install stud on both sides"):
        assert expected in text, expected
    assert "Captured but left out" not in text
    assert "ITEM 6" not in text
    # 4 photos + 1 marked-up plan sheet
    assert _picture_count(doc) == 5


def test_description_is_not_repeated_as_caption(tmp_path):
    doc = Document(str(word_report.build_report(_project(tmp_path), tmp_path / "r.docx")))
    assert _text(doc).count("Discontinued studs at window") == 2   # ITEM line + photo log


def test_report_options_turn_sections_off(tmp_path):
    project = _project(tmp_path, report_options={"include_marked_plans": False,
                                                 "include_photo_log": False})
    doc = Document(str(word_report.build_report(project, tmp_path / "r.docx")))
    text = _text(doc)
    assert "Plan Sheets" not in text and "Photo Log" not in text
    assert _picture_count(doc) == 4


def test_unnumbered_observations_get_their_own_items(tmp_path):
    project = _project(tmp_path)
    for o in project.observations:
        o.number = 0
    text = _text(Document(str(word_report.build_report(project, tmp_path / "r.docx"))))
    assert all(f"ITEM {n}\t" in text for n in (1, 2, 3))


def test_firm_template_keeps_header_and_drops_old_body(tmp_path):
    tpl = Document()
    tpl.sections[0].header.paragraphs[0].text = "ACME ENGINEERS LETTERHEAD"
    tpl.add_paragraph("Old report text that must not survive")
    tpl.save(str(tmp_path / "firm.docx"))

    project = _project(tmp_path, template_path=tmp_path / "firm.docx")
    doc = Document(str(word_report.build_report(project, tmp_path / "r.docx")))

    assert doc.sections[-1].header.paragraphs[0].text == "ACME ENGINEERS LETTERHEAD"
    assert "Old report text" not in _text(doc)
    assert "ITEM 3.1" in _text(doc)


def test_tagged_template_is_rendered_with_docxtpl(tmp_path):
    pytest.importorskip("docxtpl")
    tpl = Document()
    tpl.sections[0].header.paragraphs[0].text = "{{ project_address }} - {{ date_long }}"
    for line in ("Job {{ job_no }}: {{ project_name }}",
                 "{%p for p in structural_elements_observed %}",
                 "{{ p }}",
                 "{%p endfor %}",
                 "{%p for item in items %}",
                 "ITEM {{ item.number }} [{{ item.severity }}] {{ item.description }}",
                 "{%p for ph in item.photos %}",
                 "{{ ph }}",
                 "{%p endfor %}",
                 "{%p endfor %}"):
        tpl.add_paragraph(line)
    tpl.save(str(tmp_path / "sor.docx"))

    project = _project(tmp_path, template_path=tmp_path / "sor.docx")
    doc = Document(str(word_report.build_report(project, tmp_path / "r.docx")))
    text = _text(doc)

    assert "Job 260XX.00: 99 Green St Remodel" in text
    assert "general conformance" not in text        # nothing the engineer did not write
    assert "ITEM 3.2 [Safety] Holdown rod in (E) 4X beam" in text
    assert "{{" not in text and "{%" not in text
    assert doc.sections[0].header.paragraphs[0].text == "99 Green St, San Francisco - October 27, 2026"
    assert _picture_count(doc) == 4


def test_missing_template_fails_clearly(tmp_path):
    project = _project(tmp_path, template_path=tmp_path / "nope.docx")
    with pytest.raises(FileNotFoundError, match="template"):
        word_report.build_report(project, tmp_path / "r.docx")


def test_missing_photo_fails_clearly(tmp_path):
    project = _project(tmp_path)
    project.observations[0].photos = [tmp_path / "gone.jpg"]
    with pytest.raises(FileNotFoundError, match="gone.jpg"):
        word_report.build_report(project, tmp_path / "r.docx")


def test_same_input_gives_same_text(tmp_path):
    project = _project(tmp_path)
    a = _text(Document(str(word_report.build_report(project, tmp_path / "a.docx"))))
    b = _text(Document(str(word_report.build_report(project, tmp_path / "b.docx"))))
    assert a == b


def test_plan_sheet_pins_only_the_included_entries(tmp_path, monkeypatch):
    from sitescribe import plan_markup

    seen = []

    def fake(project, sheet_id, out):
        seen.append([o.description for o in project.observations])
        Image.new("RGB", (100, 80), "white").save(out)
        return out
    monkeypatch.setattr(plan_markup, "sheet_image_with_pins", fake)

    project = _project(tmp_path)
    project.observations[3].pin = PlanPin(project.plans[0].id, 0.1, 0.1)   # excluded, pinned
    doc = Document(str(word_report.build_report(project, tmp_path / "r.docx")))

    assert seen == [["Discontinued studs at window", "Holdown rod in (E) 4X beam", "Sheathing nailing OK"]]
    assert "Pin colors" not in _text(doc)       # plan_markup's sheet carries the legend
