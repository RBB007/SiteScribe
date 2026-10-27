"""bridge.to_project() and the Word export it feeds: capture entries through
the Flask test client, export /api/export/docx, and check the report is the
one sitescribe.word_report builds (ITEM lines, photos, pinned plan sheet)."""

import io
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from sitescribe.models import Severity
from webapp import bridge, server


@pytest.fixture
def client(tmp_path, monkeypatch):
    data, uploads = tmp_path / "data", tmp_path / "uploads"
    data.mkdir()
    uploads.mkdir()
    monkeypatch.setattr(server, "DATA_DIR", str(data))
    monkeypatch.setattr(server, "UPLOADS", str(uploads))
    monkeypatch.setattr(server, "PROJECTS_FILE", str(data / "projects.json"))
    monkeypatch.setattr(server, "LEGACY_DATA_FILE", str(data / "visit.json"))
    c = server.app.test_client()
    assert c.post("/login", data={"pin": server.PIN}).status_code == 302
    assert c.post("/api/projects", json={"name": "Elm St / Level 2"}).status_code == 200
    return c


def _jpeg(color, size=(64, 48)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "JPEG")
    buf.seek(0)
    return buf


def _plan_pdf(pages=2):
    doc = pymupdf.open()
    for _ in range(pages):
        doc.new_page(width=1200, height=800)
    return io.BytesIO(doc.tobytes())


def _item(client):
    item = client.post("/api/item").get_json()
    return item["item_id"], item["photos"][0]["photo_id"]


def _details(client, item_id, photo_id, **details):
    r = client.patch(f"/api/item/{item_id}/photo/{photo_id}", json=details)
    assert r.status_code == 200, r.get_data(as_text=True)


def _photo(client, item_id, name, color, **form):
    r = client.post(f"/api/item/{item_id}/photo", content_type="multipart/form-data",
                    data={"photo": (_jpeg(color), name), **form})
    assert r.status_code == 200, r.get_data(as_text=True)
    return r.get_json()


def _text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def _pictures(doc) -> int:
    return len(list(doc.element.body.iter(qn("wp:inline"))))


def _capture(client):
    """Two items: item 1 has two entries (one pinned on page 2 with a caption,
    one pinned on page 1 with no image yet); item 2 is left out."""
    plan = client.post("/api/plan", data={"plan": (_plan_pdf(), "Framing plans.pdf")},
                       content_type="multipart/form-data").get_json()
    beam, first_id = _item(client)
    _details(client, beam, first_id, description="Cracked beam at C-4", importance="Major",
             tags=["Crack"], location={"plan_id": plan["plan_id"], "page": 2, "x": 0.5, "y": 0.5})
    first = _photo(client, beam, "IMG_1.jpg", "red")
    _details(client, beam, first["photo_id"], caption="Looking north")
    waiting = client.post(f"/api/item/{beam}/photo", json={"inherit": "clear"}).get_json()
    _details(client, beam, waiting["photo_id"], description="Rust at base plate",
             importance="Safety", location={"plan_id": plan["plan_id"], "page": 1, "x": 0.25, "y": 0.5})

    left_out, left_out_id = _item(client)
    _details(client, left_out, left_out_id, description="Duplicate", include=False)
    _photo(client, left_out, "IMG_3.jpg", "green")
    return plan, beam, first


def test_to_project_maps_entries_plans_and_captions(client, tmp_path):
    plan, beam, first = _capture(client)
    pid = server.load_index()["projects"][0]["project_id"]
    data = server.load(pid)

    project = bridge.to_project(
        data, name="Elm St / Level 2",
        item_dir=lambda item_id: server.item_dir(item_id, pid),
        plans_dir=server.plans_dir(pid), work_dir=tmp_path / "work")

    assert project.name == "Elm St / Level 2"
    assert [o.id for o in project.observations] == [f"{beam}_01", f"{beam}_02", "002_01"]
    cracked, rust, dup = project.observations
    assert cracked.number == 1 and cracked.severity == Severity.MAJOR and cracked.tags == ["Crack"]
    assert cracked.description == "Cracked beam at C-4"          # caption not merged in
    assert cracked.captions == {Path(first["filename"]).name: "Looking north"}
    assert cracked.photos == [Path(server.item_dir(beam, pid)) / first["filename"]]
    assert cracked.pin.sheet_id == f"{plan['plan_id']}:2" and cracked.pin.x == 0.5
    assert rust.photos == [] and rust.photo_path is None
    assert rust.severity == Severity.SAFETY and rust.pin.sheet_id == f"{plan['plan_id']}:1"
    assert dup.include is False and dup.number == 2

    # both pages of the pinned plan, rendered, with the ids the pins use
    assert [s.id for s in project.plans] == [f"{plan['plan_id']}:1", f"{plan['plan_id']}:2"]
    assert project.plans[1].name == "Framing plans, page 2"
    assert all(s.image_path.is_file() for s in project.plans)
    assert project.sheet(cracked.pin.sheet_id).page_index == 1


def test_to_project_skips_unpinned_plans(client, tmp_path):
    client.post("/api/plan", data={"plan": (_plan_pdf(1), "Nobody pinned this.pdf")},
                content_type="multipart/form-data")
    item_id, photo_id = _item(client)
    _details(client, item_id, photo_id, description="Note only")
    pid = server.load_index()["projects"][0]["project_id"]

    project = bridge.to_project(
        server.load(pid), name="x",
        item_dir=lambda i: server.item_dir(i, pid),
        plans_dir=server.plans_dir(pid), work_dir=tmp_path / "work")

    assert project.plans == []
    assert not (tmp_path / "work").exists()


def test_docx_export_is_the_word_report_module(client):
    _capture(client)

    r = client.get("/api/export/docx")

    assert r.status_code == 200
    doc = Document(io.BytesIO(r.data))
    text = _text(doc)
    assert "Site Visit Report" in text and "Elm St / Level 2" in text
    assert "ITEM 1.1\tCracked beam at C-4" in text
    assert "ITEM 1.2\tRust at base plate" in text
    assert "Looking north" in text
    assert "Location:\tFraming plans, page 2" in text
    assert "Plan Sheets" in text and "Photo Log" in text
    assert "Duplicate" not in text and "ITEM 2" not in text
    # one photo + two pinned plan pages
    assert _pictures(doc) == 3


def test_docx_export_uses_the_uploaded_template(client):
    _capture(client)
    tpl = Document()
    tpl.sections[0].header.paragraphs[0].text = "ACME ENGINEERS LETTERHEAD"
    tpl.add_paragraph("Old report text that must not survive")
    buf = io.BytesIO()
    tpl.save(buf)
    buf.seek(0)
    r = client.post("/api/template", data={"template": (buf, "firm.docx")},
                    content_type="multipart/form-data")
    assert r.status_code == 200, r.get_data(as_text=True)

    doc = Document(io.BytesIO(client.get("/api/export/docx").data))

    assert doc.sections[-1].header.paragraphs[0].text == "ACME ENGINEERS LETTERHEAD"
    assert "Old report text" not in _text(doc)
    assert "ITEM 1.1" in _text(doc)


def test_docx_export_falls_back_when_the_module_fails(client, monkeypatch):
    from sitescribe import word_report

    def boom(project, out):
        raise RuntimeError("simulated module bug")
    monkeypatch.setattr(word_report, "build_report", boom)
    _capture(client)

    r = client.get("/api/export/docx")

    assert r.status_code == 200
    text = _text(Document(io.BytesIO(r.data)))
    assert "Cracked beam at C-4" in text        # the built-in report still ships
    assert "ITEM 1.1" not in text
