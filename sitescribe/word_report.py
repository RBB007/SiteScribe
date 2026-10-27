"""POST-VISIT: build the site visit report as a Word file.

Owner: Pengyu Chen (@GFDS-PC)
Library: python-docx (docxtpl only for Jinja-tagged firm templates)

Report sections, in order:
    1. Header: project, address, engineer, visit date
    2. Summary                       (project.summary)
    3. Plan sheets with color-coded pins   (plan_markup.sheet_image_with_pins)
    4. Observations: photo + caption + severity + tags, one per block
    5. Action items table            (observations with an action_item)
    6. Photo log: number, time, sheet, severity, short description

Three ways to build, picked by project.template_path (template_kind() says
which one a file gets):

* No template, or an ordinary .docx with no report in it (a letterhead, a
  blank shell): the template's body is emptied, its page setup, headers,
  footers and styles are kept, and the six sections above are written in
  with python-docx. Layout follows the GFDS Structural Observation Record:
  "ITEM 3.1 <tab> description", a "Location:" and "Action:" line, photos
  sized 2.92" tall (one portrait), 4.3" wide (one landscape) or 2.58" tall
  (two side by side).

* The firm's own report (a .docx with "ITEM n" paragraphs, e.g. a past
  Structural Observation Record): the
  document is kept as the form it is. Its header table, section labels and
  closing paragraphs stay; the date, project and other header fields are
  filled in (blank when the project has no value, never the old report's),
  the items are replaced by this visit's ITEM blocks cloned from the
  template's own paragraphs, pinned plan sheets follow them, and the old
  report's prose and previous action items are removed. See path 3 below.

* A template with Jinja tags ({{ ... }} / {%p ... %}), e.g. GFDS's
  SOR_Template.docx: rendered with docxtpl, so the template decides the
  layout. The context it gets is listed in _template_context().

Nothing is written that the engineer did not enter: no placeholder text for
a missing description, no conformance statement, no stage of construction
unless project.summary or the report option has one.

Item numbering: observations sharing a number are entries of one item.
A lone entry is "ITEM 3"; several are "ITEM 3.1", "ITEM 3.2", ...

Captions: the description is printed once, on the ITEM line. Under a photo
only its own entry in Observation.captions is printed, so the description is
not repeated under every photo.

Report options read from project.report_options (all optional):
    include_marked_plans  bool, default True   section 3
    include_photo_log     bool, default True   section 6
    job_no, site_visit_no str                  header (no Project field yet)
    stage_of_construction, structural_elements_observed
                          str                  tagged templates only; blank
                                               unless the engineer wrote them

Large phone photos are downscaled to 1600 px on the long side before they go
into the document, so a 60-photo report stays a few MB.
"""

from __future__ import annotations

import io
import logging
import re
import tempfile
import zipfile
from dataclasses import dataclass, field, replace
from pathlib import Path

from .models import SEVERITY_COLORS, Observation, PlanSheet, Project, Severity

log = logging.getLogger(__name__)

# Photo sizes, measured from GFDS sample reports (inches)
_SINGLE_PORTRAIT_H = 2.92
_SINGLE_LANDSCAPE_W = 4.30
_PAIR_H = 2.58
_PHOTO_MAX_PX = 1600
_PLAN_MAX_PX = 2400
_INDENT_IN = 1.0          # hanging indent for ITEM / Location / Action lines
_SHORT_DESC = 80          # characters of description in the photo log

_TAG_RE = re.compile(rb"\{\{|\{%")
_STEP_RE = re.compile(r"^\d+[\).]\s*")


# ----------------------------------------------------------------------------
# public
# ----------------------------------------------------------------------------
def build_report(project: Project, out_docx: Path) -> Path:
    """Write the report to out_docx and return it. If project.template_path is
    set, start from that .docx so the firm's styles, logo and header carry
    over."""
    _require("docx", "python-docx")
    out_docx = Path(out_docx)
    template = Path(project.template_path) if project.template_path else None
    if template is not None and not template.is_file():
        raise FileNotFoundError(f"Report template not found: {template}")

    entries = _entries(project)
    out_docx.parent.mkdir(parents=True, exist_ok=True)

    kind = template_kind(template) if template is not None else "styles"
    if kind == "tagged":
        _build_from_tagged_template(project, entries, template, out_docx)
    elif kind == "report":
        _build_from_firm_report(project, entries, template, out_docx)
    else:
        _build_with_python_docx(project, entries, template, out_docx)
    return out_docx


# ----------------------------------------------------------------------------
# what goes in the report
# ----------------------------------------------------------------------------
@dataclass
class _Entry:
    obs: Observation
    label: str                       # "3" or "3.1"
    photos: list[Path] = field(default_factory=list)
    sheet_name: str = ""


def _entries(project: Project) -> list[_Entry]:
    """Included observations in report order, with their ITEM labels."""
    included = [o for o in project.observations if o.include]
    # number 0 means "not numbered yet": each such entry becomes its own item
    next_free = max([o.number for o in included if o.number > 0], default=0)
    numbered: list[tuple[int, int, Observation]] = []
    for i, o in enumerate(included):
        n = o.number
        if n <= 0:
            next_free += 1
            n = next_free
        numbered.append((n, i, o))
    numbered.sort(key=lambda t: (t[0], t[1]))

    counts: dict[int, int] = {}
    for n, _, _ in numbered:
        counts[n] = counts.get(n, 0) + 1

    seen: dict[int, int] = {}
    out = []
    for n, _, o in numbered:
        seen[n] = seen.get(n, 0) + 1
        label = str(n) if counts[n] == 1 else f"{n}.{seen[n]}"
        out.append(_Entry(o, label, _photo_list(o), _sheet_name(project, o)))
    return out


def _photo_list(o: Observation) -> list[Path]:
    paths = list(o.photos) or ([o.photo_path] if o.photo_path else [])
    seen, out = set(), []
    for p in paths:
        p = Path(p)
        if p not in seen:
            seen.add(p)
            out.append(p)
    return out


def _sheet_name(project: Project, o: Observation) -> str:
    if o.pin is None:
        return ""
    sheet = next((s for s in project.plans if s.id == o.pin.sheet_id), None)
    return sheet.name if sheet else ""


def _explicit_caption(o: Observation, photo: Path) -> str:
    return (o.captions.get(Path(photo).name) or "").strip()


def _action_parts(action: str) -> tuple[str, list[str]]:
    """'Intro\\n1) step\\n2) step' -> ('Intro', ['step', 'step'])."""
    lines = [ln.strip() for ln in (action or "").splitlines() if ln.strip()]
    if len(lines) > 1 and all(_STEP_RE.match(ln) for ln in lines[1:]):
        return lines[0], [_STEP_RE.sub("", ln) for ln in lines[1:]]
    return " ".join(lines), []


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text or "") if p.strip()]


def _sev(o: Observation) -> Severity:
    return o.severity if isinstance(o.severity, Severity) else Severity(o.severity)


def _short(text: str, n: int = _SHORT_DESC) -> str:
    text = " ".join((text or "").split())
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def _time(o: Observation) -> str:
    t = o.taken_at
    return f"{t:%m/%d/%y %H:%M}" if t else ""


def _long_date(project: Project) -> str:
    d = project.visit_date
    return f"{d:%B} {d.day}, {d.year}"


# ----------------------------------------------------------------------------
# images
# ----------------------------------------------------------------------------
def _open_image(path: Path):
    from PIL import Image, ImageOps

    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Photo not found: {path}")
    if path.suffix.lower() in (".heic", ".heif"):
        try:
            import pillow_heif
        except ImportError as exc:
            raise ImportError(
                f"{path.name} is a HEIC photo: pip install pillow-heif") from exc
        pillow_heif.register_heif_opener()
    with Image.open(path) as im:
        # a transposed (or plain) copy, so the file handle closes here
        return ImageOps.exif_transpose(im)


def _image_stream(path: Path, max_px: int, *, palette: bool = False) -> tuple[io.BytesIO, int, int]:
    """Upright, downscaled copy of the image in memory, plus its size.
    palette=True writes a 256-colour PNG: a third the size for a plan sheet
    (lines, text, a few pin colours) and visually the same."""
    with _open_image(path) as im:
        im.thumbnail((max_px, max_px))
        buf = io.BytesIO()
        if palette:
            im.convert("RGB").quantize(256).save(buf, "PNG", optimize=True)
        elif im.mode in ("RGBA", "LA", "P") or path.suffix.lower() == ".png":
            im.save(buf, "PNG")
        else:
            im.convert("RGB").save(buf, "JPEG", quality=85)
        w, h = im.size
    buf.seek(0)
    return buf, w, h


def _photo_size(w: int, h: int, n: int, max_w_in: float) -> dict:
    """Width or height (inches) for one photo of a group of n."""
    from docx.shared import Inches

    if n == 1:
        if w > h:
            return {"width": Inches(min(_SINGLE_LANDSCAPE_W, max_w_in))}
        hh = _SINGLE_PORTRAIT_H
    else:
        hh = _PAIR_H
    if w / h * hh > max_w_in:
        return {"width": Inches(max_w_in)}
    return {"height": Inches(hh)}


def _plan_image(project: Project, sheet: PlanSheet, entries: list[_Entry],
                tmp: Path) -> tuple[Path | None, bool]:
    """(sheet PNG with pins, whether it already carries a legend).

    plan_markup draws the pins and a legend when it is written, else we draw
    the pins ourselves. Either way only the entries in the report are pinned:
    an observation with include=False is left off the sheet as well."""
    from . import plan_markup

    out = tmp / f"plan_{_safe_name(sheet.id)}.png"
    shown = replace(project, observations=[e.obs for e in entries])
    try:
        return Path(plan_markup.sheet_image_with_pins(shown, sheet.id, out)), True
    except NotImplementedError:
        pass
    if not sheet.image_path or not Path(sheet.image_path).is_file():
        log.warning("Plan sheet %s has no rendered image; left out of report", sheet.name)
        return None, False
    return _draw_pins(sheet, entries, out), False


def _safe_name(text: str) -> str:
    """A sheet id as a file name: "PL-01:2" -> "PL-01_2" (':' is not allowed
    on Windows)."""
    return re.sub(r"[^\w.-]", "_", str(text)) or "sheet"


def _draw_pins(sheet: PlanSheet, entries: list[_Entry], out: Path) -> Path:
    from PIL import ImageDraw, ImageFont

    with _open_image(Path(sheet.image_path)) as src:
        im = src.convert("RGB")
    im.thumbnail((_PLAN_MAX_PX, _PLAN_MAX_PX))
    w, h = im.size
    r = max(10, round(max(w, h) / 90))
    draw = ImageDraw.Draw(im)
    try:
        font = ImageFont.load_default(size=round(r * 1.1))
    except TypeError:              # Pillow < 10.1
        font = ImageFont.load_default()
    for e in entries:
        pin = e.obs.pin
        if pin is None or pin.sheet_id != sheet.id:
            continue
        x, y = pin.x * w, pin.y * h
        color = SEVERITY_COLORS[_sev(e.obs)]
        draw.ellipse([x - r, y - r, x + r, y + r], fill=color, outline="white", width=max(2, r // 6))
        draw.text((x, y), e.label, fill="white", font=font, anchor="mm")
    im.save(out, "PNG", optimize=True)
    return out


# ----------------------------------------------------------------------------
# path 1: python-docx
# ----------------------------------------------------------------------------
def _build_with_python_docx(project: Project, entries: list[_Entry],
                            template: Path | None, out_docx: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    if template is not None:
        doc = Document(str(template))
        _strip_body(doc)
    else:
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(10)

    w = _Writer(doc)
    opts = project.report_options or {}

    w.header(project)
    if project.summary.strip():
        w.heading("Summary")
        for para in _paragraphs(project.summary):
            w.body(para)
    if opts.get("include_marked_plans", True):
        with tempfile.TemporaryDirectory() as tmp:
            w.plan_sheets(project, entries, Path(tmp))
    if entries:
        w.heading("Observations")
        for e in entries:
            w.observation(e)
    w.action_items(entries)
    if opts.get("include_photo_log", True):
        w.photo_log(entries)

    doc.core_properties.title = f"Site Visit Report - {project.name}"
    _zero_image_padding(doc)
    doc.save(str(out_docx))


def _strip_body(doc) -> None:
    """Empty the template's body but keep the final <w:sectPr> (page setup,
    header/footer references), so the letterhead carries over."""
    body = doc.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)


def _zero_image_padding(doc) -> None:
    from docx.oxml.ns import qn

    for inline in doc.element.body.iter(qn("wp:inline")):
        for a in ("distT", "distB", "distL", "distR"):
            inline.set(a, "0")


class _Writer:
    """Writes the six sections. Uses the template's styles where they exist
    and falls back to direct formatting where they don't."""

    def __init__(self, doc):
        from docx.shared import Emu

        self.doc = doc
        s = doc.sections[-1]
        text_w = (s.page_width or Emu(7772400)) - (s.left_margin or Emu(914400)) \
            - (s.right_margin or Emu(914400))
        self.text_w_in = text_w / 914400

    # -- basic blocks ---------------------------------------------------
    def _styled(self, text: str, *styles: str):
        p = self.doc.add_paragraph(text)
        for name in styles:
            try:
                p.style = name
                return p
            except (KeyError, ValueError):
                continue
        return p

    def heading(self, text: str):
        from docx.shared import Pt

        p = self.doc.add_paragraph()
        try:
            p.style = "Heading 1"
        except (KeyError, ValueError):
            pass
        run = p.add_run(text)
        run.bold = True
        if p.style.name != "Heading 1":
            run.font.size = Pt(12)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        return p

    def body(self, text: str):
        from docx.shared import Pt

        p = self._styled(text, "Body Text", "Normal")
        p.paragraph_format.space_after = Pt(6)
        return p

    def _hanging(self, p):
        from docx.shared import Inches

        pf = p.paragraph_format
        pf.left_indent = Inches(_INDENT_IN)
        pf.first_line_indent = Inches(-_INDENT_IN)
        pf.tab_stops.add_tab_stop(Inches(_INDENT_IN))
        return p

    def _indented(self, p):
        from docx.shared import Inches

        p.paragraph_format.left_indent = Inches(_INDENT_IN)
        return p

    def _table(self, rows: int, cols: int, widths_in: list[float], header: list[str]):
        from docx.shared import Pt

        t = self.doc.add_table(rows=rows, cols=cols)
        try:
            t.style = "Table Grid"
        except (KeyError, ValueError):
            _set_borders(t)
        _col_widths(t, widths_in)
        for c, text in enumerate(header):
            cell = t.rows[0].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.bold = True
            _shade(cell, "D9D9D9")
        _repeat_header(t.rows[0])
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
        return t

    def _fit(self, widths: list[float]) -> list[float]:
        k = min(1.0, self.text_w_in / sum(widths))
        return [w * k for w in widths]

    # -- 1. header -------------------------------------------------------
    def header(self, project: Project):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        opts = project.report_options or {}
        p = self._styled("", "Title")
        run = p.add_run("Site Visit Report")
        if p.style.name != "Title":
            run.bold = True
            run.font.size = Pt(16)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        rows = [("Project:", project.name, "Date:", f"{project.visit_date:%m/%d/%y}"),
                ("Address:", project.site_address, "Job No:", str(opts.get("job_no", "") or "")),
                ("Engineer:", project.engineer, "Site Visit No:", str(opts.get("site_visit_no", "") or ""))]
        widths = self._fit([1.0, 3.0, 1.1, 1.4])
        t = self.doc.add_table(rows=len(rows), cols=4)
        _col_widths(t, widths)
        for r, vals in enumerate(rows):
            for c, val in enumerate(vals):
                cell = t.rows[r].cells[c]
                cell.text = ""
                run = cell.paragraphs[0].add_run(val or "")
                run.bold = c in (0, 2)
                cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        self.doc.add_paragraph()

    # -- 3. plan sheets ----------------------------------------------------
    def plan_sheets(self, project: Project, entries: list[_Entry], tmp: Path):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches

        pinned = {e.obs.pin.sheet_id for e in entries if e.obs.pin is not None}
        sheets = [s for s in project.plans if s.id in pinned]
        if not sheets:
            return
        self.heading("Plan Sheets")
        legend_shown = False
        for sheet in sheets:
            img, has_legend = _plan_image(project, sheet, entries, tmp)
            if img is None:
                continue
            if not has_legend and not legend_shown:
                # plan_markup burns a legend into the sheet; ours only has to
                # go in when we drew the pins
                self.legend()
                legend_shown = True
            cap = self.body(sheet.name)
            cap.runs[0].bold = True
            cap.paragraph_format.keep_with_next = True
            stream, w, h = _image_stream(img, _PLAN_MAX_PX, palette=True)
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            max_h = 7.5
            size = {"width": Inches(self.text_w_in)}
            if h / w * self.text_w_in > max_h:
                size = {"height": Inches(max_h)}
            p.add_run().add_picture(stream, **size)

    def legend(self):
        from docx.shared import RGBColor

        p = self.body("")
        p.add_run("Pin colors:  ")
        for sev in Severity:
            run = p.add_run("● ")
            run.font.color.rgb = RGBColor.from_string(SEVERITY_COLORS[sev].lstrip("#"))
            p.add_run(f"{sev.value}   ")

    # -- 4. observations -----------------------------------------------------
    def observation(self, e: _Entry):
        from docx.shared import Inches, Pt, RGBColor

        o = e.obs
        p = self._hanging(self.body(""))
        p.add_run(f"ITEM {e.label}").bold = True
        p.add_run("\t" + o.description.strip())
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)

        sev = _sev(o)
        p = self._indented(self.body(""))
        p.add_run("Severity: ")
        run = p.add_run(sev.value)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(SEVERITY_COLORS[sev].lstrip("#"))
        if o.tags:
            p.add_run("      Tags: " + ", ".join(o.tags))
        p.paragraph_format.keep_with_next = True

        if e.sheet_name:
            p = self._hanging(self.body(""))
            p.add_run("Location:\t" + e.sheet_name)
            p.paragraph_format.keep_with_next = True

        if e.photos:
            self.photos(e)

        intro, steps = _action_parts(o.action_item)
        if intro or steps:
            p = self._hanging(self.body(""))
            p.add_run("Action:\t" + intro)
            for i, step in enumerate(steps, 1):
                sp = self.body(f"{i})\t{step}")
                sp.paragraph_format.left_indent = Inches(_INDENT_IN + 0.25)
                sp.paragraph_format.first_line_indent = Inches(-0.25)
                sp.paragraph_format.tab_stops.add_tab_stop(Inches(_INDENT_IN + 0.25))
                sp.paragraph_format.space_after = Pt(2)

    def photos(self, e: _Entry):
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        avail = self.text_w_in - _INDENT_IN
        if len(e.photos) == 1:
            photo = e.photos[0]
            stream, w, h = _image_stream(photo, _PHOTO_MAX_PX)
            p = self._indented(self.doc.add_paragraph())
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(stream, **_photo_size(w, h, 1, avail))
            cap = _explicit_caption(e.obs, photo)
            if cap:
                # keep the photo with its caption only; chaining further would
                # glue every observation together and push them off the page
                p.paragraph_format.keep_with_next = True
                c = self._indented(self.body(cap))
                c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                c.runs[0].italic = True
            else:
                p.paragraph_format.space_after = Pt(6)
            return

        # two or more: two per row in a borderless table, caption under each
        col_w = (self.text_w_in - _INDENT_IN) / 2
        rows = (len(e.photos) + 1) // 2
        t = self.doc.add_table(rows=rows, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.RIGHT
        _col_widths(t, [col_w, col_w])
        _no_borders(t)
        for i, photo in enumerate(e.photos):
            cell = t.rows[i // 2].cells[i % 2]
            stream, w, h = _image_stream(photo, _PHOTO_MAX_PX)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.add_run().add_picture(stream, **_photo_size(w, h, 2, col_w - 0.15))
            cap = _explicit_caption(e.obs, photo)
            if cap:
                cp = cell.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.add_run(cap).italic = True
        for row in t.rows:
            _cant_split(row)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # -- 5. action items -----------------------------------------------------
    def action_items(self, entries: list[_Entry]):
        from docx.shared import RGBColor

        rows = [e for e in entries if e.obs.action_item.strip()]
        if not rows:
            return
        self.heading("Action Items")
        t = self._table(len(rows) + 1, 3, self._fit([0.8, 1.0, 4.7]), ["Item", "Severity", "Action"])
        for r, e in enumerate(rows, 1):
            cells = t.rows[r].cells
            cells[0].text = e.label
            sev = _sev(e.obs)
            cells[1].text = ""
            run = cells[1].paragraphs[0].add_run(sev.value)
            run.font.color.rgb = RGBColor.from_string(SEVERITY_COLORS[sev].lstrip("#"))
            intro, steps = _action_parts(e.obs.action_item)
            cells[2].text = intro
            for i, step in enumerate(steps, 1):
                cells[2].add_paragraph(f"{i}) {step}")

    # -- 6. photo log ----------------------------------------------------------
    def photo_log(self, entries: list[_Entry]):
        rows = [e for e in entries if e.photos]
        if not rows:
            return
        self.heading("Photo Log")
        t = self._table(len(rows) + 1, 5, self._fit([0.6, 1.25, 1.45, 0.8, 2.4]),
                        ["No.", "Time", "Sheet", "Severity", "Description"])
        for r, e in enumerate(rows, 1):
            n = len(e.photos)
            vals = [e.label + (f" ({n})" if n > 1 else ""), _time(e.obs), e.sheet_name,
                    _sev(e.obs).value, _short(e.obs.description)]
            for c, v in enumerate(vals):
                t.rows[r].cells[c].text = v


# -- table XML helpers ---------------------------------------------------------
def _col_widths(table, widths_in: list[float]) -> None:
    """Fixed column widths that Word, LibreOffice and Google Docs all honour:
    grid columns + cell widths + fixed layout."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Inches

    table.autofit = False
    tblPr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tblPr.find(qn(tag))
        if old is not None:
            tblPr.remove(old)
    total = sum(round(w * 1440) for w in widths_in)
    tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{total}" w:type="dxa"/>'))
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths_in):
        gc.set(qn("w:w"), str(round(w * 1440)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in):
            cell.width = Inches(w)


def _borders_xml(val: str) -> str:
    from docx.oxml.ns import nsdecls

    edges = "".join(f'<w:{k} w:val="{val}" w:sz="4" w:space="0" w:color="808080"/>'
                    for k in ("top", "left", "bottom", "right", "insideH", "insideV"))
    return f"<w:tblBorders {nsdecls('w')}>{edges}</w:tblBorders>"


def _put_borders(table, val: str) -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    tblPr = table._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(parse_xml(_borders_xml(val)))


def _set_borders(table) -> None:
    _put_borders(table, "single")


def _no_borders(table) -> None:
    _put_borders(table, "nil")


def _shade(cell, hex_fill: str) -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_fill}"/>'))


def _repeat_header(row) -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def _cant_split(row) -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


# ----------------------------------------------------------------------------
# path 2: Jinja-tagged template (docxtpl)
# ----------------------------------------------------------------------------
def _has_jinja_tags(template: Path) -> bool:
    with zipfile.ZipFile(template) as z:
        for name in z.namelist():
            if name.startswith("word/") and name.endswith(".xml") and (
                    "document" in name or "header" in name or "footer" in name):
                if _TAG_RE.search(z.read(name)):
                    return True
    return False


def _build_from_tagged_template(project: Project, entries: list[_Entry],
                                template: Path, out_docx: Path) -> None:
    _require("docxtpl", "docxtpl")
    from docxtpl import DocxTemplate

    tpl = DocxTemplate(str(template))
    with tempfile.TemporaryDirectory() as tmp:
        ctx = _template_context(tpl, project, entries, Path(tmp))
        tpl.render(ctx, autoescape=True)
    _zero_image_padding(tpl.docx)
    tpl.docx.core_properties.title = f"Site Visit Report - {project.name}"
    tpl.save(str(out_docx))


def _template_context(tpl, project: Project, entries: list[_Entry], tmp: Path) -> dict:
    """Variables a tagged template can use.

    GFDS SOR_Template.docx names (unchanged from SOR_Report_App):
        date_short, date_long, job_no, project_name, project_address,
        observer, site_visit_no, stage_of_construction [paragraphs],
        structural_elements_observed [paragraphs], previous_items [],
        items [{number, description, location, caption, photos [images],
                action, action_list [str], photos_after_action []}]
    Extra, for templates that want them:
        summary [paragraphs], legend [{name, color}],
        items[].severity, .severity_color, .tags, .time, .captions [str],
        action_items [{number, severity, action, action_list}],
        photo_log [{number, time, sheet, severity, description}],
        plans [{name, image}]
    """
    from docx.shared import Inches
    from docxtpl import InlineImage

    opts = project.report_options or {}

    def images(e: _Entry) -> list:
        # SOR template: a group sits on one line right of a 1.5" indent (~4.4"
        # usable) with ~0.55" gaps between photos
        n = len(e.photos)
        max_w = 4.3 if n <= 1 else max(1.0, (4.2 - 0.6 * (n - 1)) / n)
        out = []
        for p in e.photos:
            stream, w, h = _image_stream(p, _PHOTO_MAX_PX)
            out.append(InlineImage(tpl, stream, **_photo_size(w, h, n, max_w)))
        return out

    items = []
    for e in entries:
        o = e.obs
        caps = [_explicit_caption(o, p) for p in e.photos]
        distinct = [c for c in dict.fromkeys(caps) if c]
        action, steps = _action_parts(o.action_item)
        sev = _sev(o)
        items.append({
            "number": e.label,
            "description": o.description.strip(),
            "location": e.sheet_name,
            "caption": "; ".join(distinct),
            "captions": caps,
            "photos": images(e),
            "action": action,
            "action_list": steps,
            "photos_after_action": [],
            "severity": sev.value,
            "severity_color": SEVERITY_COLORS[sev],
            "tags": list(o.tags),
            "time": _time(o),
        })

    plans = []
    if opts.get("include_marked_plans", True):
        pinned = {e.obs.pin.sheet_id for e in entries if e.obs.pin is not None}
        for sheet in (s for s in project.plans if s.id in pinned):
            img, _ = _plan_image(project, sheet, entries, tmp)
            if img is not None:
                stream, _, _ = _image_stream(img, _PLAN_MAX_PX, palette=True)
                plans.append({"name": sheet.name, "image": InlineImage(tpl, stream, width=Inches(6.5))})

    # Nothing is written into the record that the engineer did not enter:
    # the stage and the conformance statement stay blank unless given.
    summary = _paragraphs(project.summary)
    return {
        "date_short": f"{project.visit_date:%m/%d/%y}",
        "date_long": _long_date(project),
        "job_no": str(opts.get("job_no", "") or ""),
        "project_name": project.name,
        "project_address": project.site_address,
        "observer": project.engineer,
        "site_visit_no": str(opts.get("site_visit_no", "") or ""),
        "stage_of_construction": _paragraphs(opts.get("stage_of_construction") or ""),
        "structural_elements_observed": _paragraphs(opts.get("structural_elements_observed") or ""),
        "summary": summary,
        "items": items,
        "previous_items": list(opts.get("previous_items") or []),
        "action_items": [{"number": i["number"], "severity": i["severity"],
                          "action": i["action"], "action_list": i["action_list"]}
                         for i in items if i["action"] or i["action_list"]],
        "photo_log": [{"number": e.label, "time": _time(e.obs), "sheet": e.sheet_name,
                       "severity": _sev(e.obs).value, "description": _short(e.obs.description)}
                      for e in entries if e.photos],
        "plans": plans,
        "legend": [{"name": s.value, "color": SEVERITY_COLORS[s]} for s in Severity],
    }


# ----------------------------------------------------------------------------
# path 3: the firm's own report used as the template (no Jinja tags)
# ----------------------------------------------------------------------------
# A past Structural Observation Record uploaded as the template is treated as
# a form: its header table, "Stage of Construction", "Structural Elements
# Observed", "Specific Items Observed" heading, "Previous Action Items" and
# the closing paragraphs ("Action items in this report...", "End of Report")
# are kept where they are, and only the parts that change from visit to
# visit are rewritten:
#
#     Date / Project / Job No / Site Visit No / By / Address
#                                      set from the project; blank when the
#                                      project does not have them. Nothing
#                                      from the old report is carried over.
#     running headers and footers      dates become the visit date; in a
#                                      header that carries a date, the plain
#                                      text lines (the old project line)
#                                      become the project name
#     Stage of Construction            project.summary (or the report option),
#                                      else left blank
#     Structural Elements Observed     the report option, else left blank
#     Specific Items Observed          rebuilt: one ITEM block per entry,
#                                      cloned from the template's own ITEM,
#                                      photo and Action: paragraphs
#     Plan Sheets                      added after the items when any entry
#                                      is pinned (include_marked_plans)
#     Previous Action Items            report option previous_items, else the
#                                      section is removed
_ITEM_RE = re.compile(r"^\s*ITEM\s+\d", re.I)
_ACTION_RE = re.compile(r"^\s*Action\s*:", re.I)
_STATUS_RE = re.compile(r"^\s*Status\s*:", re.I)
_BLOCK_LABEL_RE = re.compile(r"^\s*(ITEM\s+\d|Action\s*:|Status\s*:|Location\s*:|Severity\s*:)", re.I)
_ITEMS_HEADING_RE = re.compile(r"^\s*Specific\s+Items\s+Observed", re.I)
_PREVIOUS_RE = re.compile(r"^\s*Previous\s+Action\s+Items", re.I)
_STAGE_RE = re.compile(r"^\s*Stage\s+of\s+Construction", re.I)
_ELEMENTS_RE = re.compile(r"^\s*Structural\s+Elements\s+Observed", re.I)
_LONG_DATE_RE = re.compile(
    r"\b(?:January|February|March|April|May|June|July|August|September|October|"
    r"November|December)\s+\d{1,2}\s*,\s*\d{4}\b")
_SHORT_DATE_RE = re.compile(r"\b\d{1,2}/\d{1,2}/(?:\d{4}|\d{2})\b")
_HEADER_FIELDS = (
    # label pattern, project value. Every field is written, blank when the
    # project has no value: an old report's job number or observer must not
    # appear on this one.
    (r"Date", lambda p, o: f"{p.visit_date:%m/%d/%y}"),
    (r"Project(?:\s+Name)?", lambda p, o: p.name),
    (r"Site\s*Visit\s*(?:No|Number|#)\.?", lambda p, o: str(o.get("site_visit_no", "") or "")),
    (r"Job\s*(?:No|Number|#)\.?", lambda p, o: str(o.get("job_no", "") or "")),
    (r"By|Observer|Engineer|Observed\s+By", lambda p, o: p.engineer),
    (r"(?:Site\s+|Project\s+)?Address", lambda p, o: p.site_address),
)


def template_kind(template: Path) -> str:
    """How build_report will use a template: "tagged" (Jinja tags, rendered
    with docxtpl), "report" (a firm report with "ITEM n" paragraphs: kept and
    refilled, the ITEM blocks cloned from those paragraphs) or "styles" (any
    other .docx, a form with headings but no items included: body replaced,
    page setup and styles kept)."""
    _require("docx", "python-docx")
    from docx import Document

    template = Path(template)
    if _has_jinja_tags(template):
        return "tagged"
    return "report" if _is_firm_report(Document(str(template))) else "styles"


def _is_firm_report(doc) -> bool:
    """True when the body has an "ITEM n" paragraph: the one thing the form
    path cannot do without, since every new block is cloned from it."""
    return any(_ITEM_RE.match(_el_text(el)) for el in _body_elements(doc))


def _build_from_firm_report(project: Project, entries: list[_Entry],
                            template: Path, out_docx: Path) -> None:
    from docx import Document

    doc = Document(str(template))
    opts = project.report_options or {}
    els = _body_elements(doc)
    start, end = _items_region(els)
    protos = _Protos.pick(els[start:end], heading=els[start - 1] if start else None)

    # header table, running headers and footers
    _fill_header_fields(doc, els[:start], project)
    _fill_dates(doc, els[:start], project)

    # the two prose sections above the items
    stage = _paragraphs(opts.get("stage_of_construction") or project.summary)
    _fill_section(doc, els, _STAGE_RE, start, stage, clear=True)
    elements = _paragraphs(opts.get("structural_elements_observed") or "")
    _fill_section(doc, els, _ELEMENTS_RE, start, elements, clear=True)

    # previous action items: rewrite from the option, or drop the section
    prev = _previous_region(els, end)
    previous = list(opts.get("previous_items") or [])
    if prev is not None:
        p_start, p_end = prev
        if previous:
            _rewrite_previous(doc, els[p_start], els[p_start + 1:p_end], previous, protos)
        else:
            _remove(els[p_start:p_end])

    # the items themselves, written before whatever follows them that is
    # still in the document (the closing text, or the end of the body)
    _remove(els[start:end])
    anchor = next((el for el in els[end:] if el.getparent() is not None),
                  doc.element.body.sectPr)
    writer = _FormWriter(doc, protos, anchor)
    for e in entries:
        writer.observation(e)
    if opts.get("include_marked_plans", True):
        with tempfile.TemporaryDirectory() as tmp:
            writer.plan_sheets(project, entries, Path(tmp))
    writer.blank()

    doc.core_properties.title = f"Site Visit Report - {project.name}"
    _zero_image_padding(doc)
    doc.save(str(out_docx))


# -- reading the template ------------------------------------------------------
def _body_elements(doc) -> list:
    """Top-level paragraphs and tables of the body, in order (no sectPr)."""
    from docx.oxml.ns import qn

    return [el for el in doc.element.body.iterchildren()
            if el.tag in (qn("w:p"), qn("w:tbl"))]


def _el_text(el) -> str:
    """Text of a paragraph element (tabs as \\t); "" for a table."""
    from docx.oxml.ns import qn

    if el.tag != qn("w:p"):
        return ""
    parts = []
    for node in el.iter():
        if node.getparent().tag != qn("w:r"):
            continue        # a <w:tab> under <w:pPr><w:tabs> is a tab stop, not text
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag == qn("w:tab"):
            parts.append("\t")
    return "".join(parts)


def _has_picture(el) -> bool:
    return bool(el.xpath('.//*[local-name()="blip" or local-name()="imagedata"]'))


def _left_indent(el) -> int:
    """Left indent in twips, 0 when none."""
    from docx.oxml.ns import qn

    pPr = el.find(qn("w:pPr"))
    ind = pPr.find(qn("w:ind")) if pPr is not None else None
    if ind is None:
        return 0
    try:
        return int(ind.get(qn("w:left")) or ind.get(qn("w:start")) or 0)
    except ValueError:
        return 0


def _is_numbered(el) -> bool:
    from docx.oxml.ns import qn

    pPr = el.find(qn("w:pPr"))
    return pPr is not None and pPr.find(qn("w:numPr")) is not None


def _is_block_content(el) -> bool:
    """Part of an ITEM block (or a spacer), as opposed to a section heading or
    a closing paragraph: empty, indented, numbered, a picture, a table, or a
    line that starts with ITEM / Action: / Status: / Location: / Severity:."""
    from docx.oxml.ns import qn

    if el.tag != qn("w:p"):
        return True
    text = _el_text(el).strip()
    return (not text or _has_picture(el) or _left_indent(el) > 0 or _is_numbered(el)
            or bool(_BLOCK_LABEL_RE.match(text)))


def _items_region(els: list) -> tuple[int, int]:
    """[start, end) of the ITEM blocks: everything after the "Specific Items
    Observed" heading (or from the first ITEM paragraph) up to the first
    paragraph after the last ITEM that is not block content (Previous Action
    Items, the closing text, End of Report). Scanning from the last ITEM, not
    the first, means an unindented line inside the old items (a sub-heading,
    a description that wrapped into its own paragraph) does not end the
    region early and leave the old report's later items in the document."""
    start = None
    for i, el in enumerate(els):
        if _ITEMS_HEADING_RE.match(_el_text(el)):
            start = i + 1
            break
    items = [i for i, el in enumerate(els) if _ITEM_RE.match(_el_text(el))]
    if start is None:
        start = items[0]
    # "Previous Action Items" has ITEM lines of its own: the region never
    # reaches past its heading
    limit = next((i for i in range(start, len(els)) if _PREVIOUS_RE.match(_el_text(els[i]))), len(els))
    last = max([i for i in items if start <= i < limit], default=start)
    end = limit
    for i in range(last, limit):
        if not _is_block_content(els[i]):
            end = i
            break
    return start, end


def _previous_region(els: list, items_end: int) -> tuple[int, int] | None:
    """[start, end) of the "Previous Action Items" section, heading included,
    when it follows the items."""
    if items_end >= len(els) or not _PREVIOUS_RE.match(_el_text(els[items_end])):
        return None
    end = len(els)
    for i in range(items_end + 1, len(els)):
        if not _is_block_content(els[i]):
            end = i
            break
    return items_end, end


def _remove(els) -> None:
    for el in els:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


@dataclass
class _Protos:
    """Paragraphs of the template the new blocks are cloned from, so fonts,
    indents and tab stops are the firm's, not ours."""
    item: object
    action: object
    photo: object
    status: object | None
    heading: object | None
    indent_in: float          # left indent of an ITEM block, inches

    @classmethod
    def pick(cls, region: list, heading) -> "_Protos":
        from docx.oxml.ns import qn

        paras = [el for el in region if el.tag == qn("w:p")]
        items = [el for el in paras if _ITEM_RE.match(_el_text(el))]
        if not items:
            raise ValueError("The report template has no 'ITEM n' paragraph to copy the layout from")
        item = next((el for el in items if _left_indent(el) > 0), items[0])
        action = next((el for el in paras if _ACTION_RE.match(_el_text(el))), item)
        photo = next((el for el in paras if _has_picture(el)), None)
        status = next((el for el in paras if _STATUS_RE.match(_el_text(el))), None)
        indent = _left_indent(item) / 1440 or _INDENT_IN
        return cls(item, action, photo if photo is not None else action, status, heading, indent)


# -- filling the parts that stay -----------------------------------------------
def _fill_header_fields(doc, head: list, project: Project) -> None:
    """Date:, Project:, Job No: ... in the tables (and plain "Label: value"
    paragraphs) above the items. A label alone in a cell names the cell to
    its right; "Label: value" in one paragraph is rewritten in place."""
    from docx.oxml.ns import qn
    from docx.table import Table

    opts = project.report_options or {}
    for el in head:
        if el.tag == qn("w:tbl"):
            for row in Table(el, doc._body).rows:
                cells = []
                for c in row.cells:
                    if not cells or c._tc is not cells[-1]._tc:
                        cells.append(c)
                for i, cell in enumerate(cells):
                    nxt = cells[i + 1] if i + 1 < len(cells) else None
                    for p in cell.paragraphs:
                        _fill_field(p, nxt, project, opts)
        else:
            _fill_field(_para(doc, el), None, project, opts)


def _fill_field(p, next_cell, project: Project, opts: dict) -> None:
    text = _el_text(p._p)      # positions as _replace_span counts them, not Paragraph.text's
    for label, value_of in _HEADER_FIELDS:
        m = re.match(rf"^\s*(?:{label})\s*:(.*)$", text, re.I | re.S)
        if not m:
            continue
        value = (value_of(project, opts) or "").strip()
        if m.group(1).strip():
            # "Site Visit No:     2" -> keep the label and the gap, swap the value
            gap = len(m.group(1)) - len(m.group(1).lstrip())
            _replace_span(p, m.start(1) + gap, len(text), value)
        elif next_cell is not None:
            _set_text(next_cell.paragraphs[0], value)
        elif value:
            _replace_span(p, len(text), len(text), " " + value)
        return


def _fill_dates(doc, head: list, project: Project) -> None:
    """Dates in the running headers and footers become the visit date, long
    or short to match what was there. Above the items only a date that
    stands on its own line (or after a "Date:" label) is the report date and
    is swapped; a date inside other text ("Drawings dated 3/4/25") is left as
    it is, since it may well not be this visit's.

    A running header that carries the report date is per-report, so the
    plain text line next to the date (no fields, no pictures: the old
    report's project line, the nearest one above the date, else the nearest
    below) becomes this project's name. Any other plain line in that header,
    and a header with no date at all, e.g. just the firm's name, is left
    alone."""
    from docx.oxml.ns import qn

    long_date, short_date = _long_date(project), f"{project.visit_date:%m/%d/%y}"

    def dates_in(p):
        text = _el_text(p._p)      # positions as _replace_span counts them
        return text, [(m, value) for pattern, value in ((_LONG_DATE_RE, long_date),
                                                        (_SHORT_DATE_RE, short_date))
                      for m in pattern.finditer(text)]

    def swap_dates(p, *, alone: bool = False) -> bool:
        text, found = dates_in(p)
        if not found:
            return False
        if alone:
            rest = text
            for m, _ in found:
                rest = rest.replace(m.group(), "")
            if rest.strip() and not re.match(r"^\s*Date\s*:\s*$", rest, re.I):
                return False
        for m, value in sorted(found, key=lambda mv: -mv[0].start()):
            _replace_span(p, m.start(), m.end(), value)
        return True

    def plain(p) -> bool:
        return bool(_el_text(p._p).strip()) and not _has_picture(p._p) \
            and not p._p.xpath('.//w:fldChar | .//w:fldSimple | .//w:instrText')

    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header,
                     section.even_page_footer):
            if part.is_linked_to_previous:
                continue
            paras = [_para(doc, el) for el in part._element.iter(qn("w:p"))]
            dated = [swap_dates(p) for p in paras]
            if not any(dated) or part._element.tag != qn("w:hdr"):
                continue
            first = dated.index(True)
            above = [p for p in paras[:first] if plain(p)]
            below = [p for p, d in zip(paras[first + 1:], dated[first + 1:]) if plain(p) and not d]
            line = above[-1] if above else below[0] if below else None
            if line is not None:
                _set_text(line, project.name)
    for el in head:
        for p_el in el.iter(qn("w:p")):
            swap_dates(_para(doc, p_el), alone=True)


def _fill_section(doc, els: list, label_re, stop: int, texts: list[str], *, clear: bool) -> None:
    """The paragraphs under a "Stage of Construction:"-style label, up to the
    next label (a paragraph with text and no indent) or the items, become
    `texts`; the first one is the prototype for the rest. With no texts the
    section is left with one empty line."""
    from docx.oxml.ns import qn

    start = next((i for i in range(stop) if label_re.match(_el_text(els[i]))), None)
    if start is None:
        return
    body = []
    for i in range(start + 1, stop):
        el = els[i]
        if el.tag != qn("w:p"):
            break
        if _el_text(el).strip() and _left_indent(el) == 0 and not _is_numbered(el):
            break
        body.append(el)
    filled = [el for el in body if _el_text(el).strip()]
    proto = filled[0] if filled else None
    if proto is None:
        # a blank form: write after the label with the label's indent and the body font
        proto = _clone(els[start], keep_runs=False)
        els[start].addnext(proto)
        _para(doc, proto).paragraph_format.left_indent = None
        for r in proto.iter(qn("w:rPr")):
            for tag in ("w:b", "w:bCs"):
                b = r.find(qn(tag))
                if b is not None:
                    r.remove(b)
        filled, body = [proto], body + [proto]
    if clear:
        for el in filled[1:]:
            _remove([el])
    if not texts:
        _set_text(_para(doc, proto), "")
        return
    _set_text(_para(doc, proto), texts[0])
    last = proto
    for text in texts[1:]:
        new = _clone(proto, keep_runs=True)
        last.addnext(new)
        _set_text(_para(doc, new), text)
        last = new


def _rewrite_previous(doc, heading, old: list, previous: list, protos: _Protos) -> None:
    """ITEM n / Status: lines under "Previous Action Items:" from the option:
    strings ("ITEM 1.2: Closed") or {number, status, description} dicts."""
    _remove(old)
    anchor = heading.getnext()
    body = doc._body
    status_proto = protos.status if protos.status is not None else protos.action
    for entry in previous:
        if isinstance(entry, dict):
            number = str(entry.get("number", "")).strip()
            text = str(entry.get("description", "")).strip()
            status = str(entry.get("status", "")).strip()
        else:
            number, _, status = str(entry).partition(":")
            number, status, text = number.replace("ITEM", "").strip(), status.strip(), ""
        _insert_clone(protos.item, anchor, body, [(f"ITEM {number}", True), ("\t" + text, None)])
        _insert_clone(protos.item, anchor, body, [])
        _insert_clone(status_proto, anchor, body,
                      [("Status: " if protos.status is not None else "Status:\t", None), (status, None)])
        _insert_clone(protos.item, anchor, body, [])
    _insert_clone(protos.item, anchor, body, [])


# -- writing the new blocks ----------------------------------------------------
class _FormWriter:
    """Writes ITEM blocks (and the plan sheets) before `anchor`, each
    paragraph a clone of the template's own."""

    def __init__(self, doc, protos: _Protos, anchor):
        from docx.shared import Emu

        self.doc, self.protos, self.anchor = doc, protos, anchor
        s = doc.sections[-1]
        text_w = (s.page_width or Emu(7772400)) - (s.left_margin or Emu(914400)) \
            - (s.right_margin or Emu(914400))
        self.text_w_in = text_w / 914400
        self.avail_in = max(2.0, self.text_w_in - protos.indent_in)

    def _add(self, proto, segments):
        return _insert_clone(proto, self.anchor, self.doc._body, segments)

    def blank(self):
        return self._add(self.protos.item, [])

    def observation(self, e: _Entry):
        from docx.shared import Inches, Pt, RGBColor

        o = e.obs
        self.blank()
        self._add(self.protos.item, [(f"ITEM {e.label}", True),
                                     ("\t" + o.description.strip(), None)])
        sev = _sev(o)
        p = self._add(self.protos.action, [("Severity:\t", None)])
        run = p.add_run(sev.value)
        _copy_rpr(p, run)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(SEVERITY_COLORS[sev].lstrip("#"))
        if o.tags:
            _copy_rpr(p, p.add_run("      Tags: " + ", ".join(o.tags)))
        if e.sheet_name:
            self._add(self.protos.action, [("Location:\t" + e.sheet_name, None)])
        if e.photos:
            self.blank()
            self.photos(e)
        intro, steps = _action_parts(o.action_item)
        if intro or steps:
            self.blank()
            self._add(self.protos.action, [("Action:\t" + intro, None)])
            for i, step in enumerate(steps, 1):
                sp = self._add(self.protos.action, [(f"{i})\t{step}", None)])
                pf = sp.paragraph_format
                pf.left_indent = Inches(self.protos.indent_in + 0.25)
                pf.first_line_indent = Inches(-0.25)
                pf.tab_stops.add_tab_stop(Inches(self.protos.indent_in + 0.25))
                pf.space_after = Pt(2)
        self.blank()

    def _picture_para(self, indent_in: float):
        from docx.shared import Inches

        el = _clone(self.protos.photo, keep_runs=False)
        self.anchor.addprevious(el)
        p = _para(self.doc, el)
        pf = p.paragraph_format
        pf.left_indent = Inches(indent_in)
        pf.first_line_indent = Inches(0)
        return p

    def photos(self, e: _Entry):
        from docx.shared import Pt

        photos = e.photos
        if len(photos) == 1:
            stream, w, h = _image_stream(photos[0], _PHOTO_MAX_PX)
            p = self._picture_para(self.protos.indent_in)
            p.add_run().add_picture(stream, **_photo_size(w, h, 1, self.avail_in))
            self._caption(e, photos[0])
            return
        # two per line, as in the firm's reports, each caption on the line below
        col_w = (self.avail_in - 0.3) / 2
        for i in range(0, len(photos), 2):
            pair = photos[i:i + 2]
            p = self._picture_para(self.protos.indent_in)
            for j, photo in enumerate(pair):
                if j:
                    _copy_rpr(p, p.add_run("   "))
                stream, w, h = _image_stream(photo, _PHOTO_MAX_PX)
                p.add_run().add_picture(stream, **_photo_size(w, h, 2, col_w))
            caps = [_explicit_caption(e.obs, ph) for ph in pair]
            if any(caps):
                c = self._add(self.protos.action, [])
                c.paragraph_format.first_line_indent = Pt(0)
                for j, cap in enumerate(caps):
                    if j:
                        _copy_rpr(c, c.add_run("\t"))
                    run = c.add_run(cap)
                    _copy_rpr(c, run)
                    run.italic = True

    def _caption(self, e: _Entry, photo: Path):
        from docx.shared import Pt

        cap = _explicit_caption(e.obs, photo)
        if not cap:
            return
        p = self._add(self.protos.action, [])
        p.paragraph_format.first_line_indent = Pt(0)
        run = p.add_run(cap)
        _copy_rpr(p, run)
        run.italic = True

    def plan_sheets(self, project: Project, entries: list[_Entry], tmp: Path):
        from docx.shared import Inches

        pinned = {e.obs.pin.sheet_id for e in entries if e.obs.pin is not None}
        sheets = [s for s in project.plans if s.id in pinned]
        if not sheets:
            return
        heading = self.protos.heading if self.protos.heading is not None else self.protos.action
        lead = re.match(r"\s*", _el_text(heading)).group() if self.protos.heading is not None else ""
        self.blank()
        self._add(heading, [(lead + "Plan Sheets", True)])
        for sheet in sheets:
            img, _ = _plan_image(project, sheet, entries, tmp)
            if img is None:
                continue
            self.blank()
            name = self._add(self.protos.action, [(sheet.name, True)])
            pf = name.paragraph_format
            pf.left_indent = Inches(0)
            pf.first_line_indent = Inches(0)
            pf.keep_with_next = True
            stream, w, h = _image_stream(img, _PLAN_MAX_PX, palette=True)
            size = {"width": Inches(self.text_w_in)}
            if h / w * self.text_w_in > 7.0:
                size = {"height": Inches(7.0)}
            self._picture_para(0).add_run().add_picture(stream, **size)


# -- paragraph XML helpers -------------------------------------------------------
def _para(doc, el):
    from docx.text.paragraph import Paragraph

    return Paragraph(el, doc._body)


def _clone(proto, *, keep_runs: bool):
    """A copy of a paragraph element: its pPr, and its runs only if asked."""
    from copy import deepcopy

    from docx.oxml.ns import qn

    el = deepcopy(proto)
    for child in list(el):
        if child.tag == qn("w:pPr"):
            continue
        if keep_runs and child.tag == qn("w:r"):
            continue
        el.remove(child)
    return el


def _base_rpr(proto):
    """Run properties for body text copied from a paragraph: the first
    non-bold run's, else the first run's with bold removed; None when the
    paragraph has no runs (the style then decides)."""
    from copy import deepcopy

    from docx.oxml.ns import qn

    runs = proto.findall(qn("w:r"))
    if not runs:
        return None
    plain = next((r for r in runs if r.find(qn("w:rPr")) is None
                  or r.find(qn("w:rPr")).find(qn("w:b")) is None), runs[0])
    rpr = plain.find(qn("w:rPr"))
    if rpr is None:
        return None
    rpr = deepcopy(rpr)
    for tag in ("w:b", "w:bCs", "w:i", "w:iCs", "w:color", "w:u"):
        for node in rpr.findall(qn(tag)):
            rpr.remove(node)
    return rpr


def _copy_rpr(paragraph, run) -> None:
    """Give a run added to a cloned paragraph that paragraph's body font."""
    from copy import deepcopy

    rpr = getattr(paragraph._p, "_sitescribe_rpr", None)
    if rpr is not None:
        run._r.insert(0, deepcopy(rpr))


def _insert_clone(proto, anchor, body, segments):
    """Clone `proto` before `anchor` and write `segments` [(text, bold)]
    into it in the prototype's body font. Returns the Paragraph."""
    from copy import deepcopy

    from docx.text.paragraph import Paragraph

    el = _clone(proto, keep_runs=False)
    anchor.addprevious(el)
    p = Paragraph(el, body)
    rpr = _base_rpr(proto)
    el._sitescribe_rpr = rpr      # for runs added later (severity colour, captions)
    for text, bold in segments:
        if not text:
            continue
        run = p.add_run(text)
        if rpr is not None:
            run._r.insert(0, deepcopy(rpr))
        if bold is not None:
            run.bold = bold
    return p


def _set_text(p, text: str) -> None:
    """Replace a paragraph's text, keeping the first run's formatting."""
    from docx.oxml.ns import qn

    runs = p._p.findall(qn("w:r"))
    for r in runs[1:]:
        p._p.remove(r)
    if runs:
        for child in list(runs[0]):
            if child.tag != qn("w:rPr"):
                runs[0].remove(child)
        run = p.runs[0]
        run.text = text
    else:
        p.add_run(text)


def _replace_span(p, start: int, end: int, text: str) -> None:
    """Replace characters [start, end) of a paragraph's text (as _el_text
    counts it, tabs included) with `text`, across however many runs the span
    covers, keeping each run's formatting."""
    from docx.oxml.ns import qn

    pos, first, done = 0, True, False
    for node in list(p._p.iter()):
        if node.getparent().tag != qn("w:r"):
            continue
        if node.tag == qn("w:tab"):
            pos += 1
            continue
        if node.tag != qn("w:t"):
            continue
        t = node.text or ""
        t_start, t_end = pos, pos + len(t)
        pos = t_end
        if t_end < start or t_start > end or (t_start == end and not first):
            continue
        if t_end == start and not (first and start == end):
            continue
        head = t[:max(0, start - t_start)]
        tail = t[max(0, end - t_start):]
        node.text = head + (text if first else "") + tail
        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        first = False
        done = True
    if not done and first:
        p.add_run(text)


# ----------------------------------------------------------------------------
def _require(module: str, package: str) -> None:
    try:
        __import__(module)
    except ImportError as exc:
        raise ImportError(f"word_report needs {package}: pip install {package}") from exc
