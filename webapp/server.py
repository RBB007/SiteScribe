"""SiteScribe — structural site-visit observation tool.

Single-file Flask backend. Projects hold a visit; a visit holds items;
an item is only a numbering group for photo entries, and a photo *entry*
is the observation itself — description, importance, tags, plan location,
audio clips, include flag and (optionally) an image.

Ids: item "001", photo entry "001_01", audio clip "001_01_A01". An entry
whose filename is None is a note taken before the picture.

Storage:
    data/projects.json            index of every project
    data/<pid>/visit.json         that project's items + plans
    data/<pid>/template.docx      that project's Word report template, if any
    uploads/<pid>/<item_id>/...   that project's photos + audio
    uploads/<pid>/plans/...       that project's plan files

The open project lives in session["pid"], so the item/photo/audio/plan
routes keep their original URLs and resolve their paths against it.
"""

from flask import (Flask, render_template, render_template_string, request,
                   jsonify, session, redirect, url_for, send_file,
                   send_from_directory)
from werkzeug.exceptions import HTTPException
from werkzeug.utils import secure_filename

import datetime
import functools
import io
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import traceback
from pathlib import Path

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE, "data")
UPLOADS = os.path.join(BASE, "uploads")
PROJECTS_FILE = os.path.join(DATA_DIR, "projects.json")
LEGACY_DATA_FILE = os.path.join(DATA_DIR, "visit.json")

for d in (DATA_DIR, UPLOADS):
    os.makedirs(d, exist_ok=True)

IMPORTANCE = ["Info", "Minor", "Major", "Critical", "Safety"]
PID_RE = re.compile(r"^P-\d{3}$")
MAX_NAME = 120
MAX_ITEMS = 999      # item ids are 3 digits
MAX_PHOTOS = 99      # photo seq is 2 digits
INHERIT_MODES = ("keep", "clear", "new")

app = Flask(__name__)
app.secret_key = os.environ.get("SITESCRIBE_SECRET", os.urandom(32).hex())
app.permanent_session_lifetime = datetime.timedelta(hours=12)
app.config["MAX_CONTENT_LENGTH"] = 64 * 1024 * 1024

PIN = os.environ.get("SITESCRIBE_PIN", "2424")

# ---------------------------------------------------------------- auth

# crude per-IP throttle: a 4-digit PIN is only 10k guesses
_fails = {}
MAX_FAILS, LOCK_SECONDS = 8, 300


def _ip():
    fwd = request.headers.get("CF-Connecting-IP") or request.headers.get("X-Forwarded-For")
    return (fwd.split(",")[0].strip() if fwd else request.remote_addr) or "?"


def locked(ip):
    n, until = _fails.get(ip, (0, 0))
    return time.time() < until


def require_pin(fn):
    @functools.wraps(fn)
    def wrapper(*a, **kw):
        if not session.get("ok"):
            if request.path.startswith("/api/"):
                return jsonify(error="unauthorized"), 401
            return redirect(url_for("login", next=request.path))
        return fn(*a, **kw)
    return wrapper


LOGIN_HTML = """<!DOCTYPE html><html lang=en><head><meta charset=utf-8>
<meta name=viewport content="width=device-width,initial-scale=1">
<title>SiteScribe</title><style>
:root{--bg:#fff;--fg:#14181d;--mut:#667085;--line:#e4e7ec;--card:#f9fafb;--accent:#1554b8;--bad:#b42318}
@media(prefers-color-scheme:dark){:root{--bg:#101418;--fg:#e8eaed;--mut:#98a2b3;--line:#272c33;--card:#171b20;--accent:#4d8df0;--bad:#ff8a80}}
*{box-sizing:border-box}body{margin:0;min-height:100vh;display:grid;place-items:center;
background:var(--bg);color:var(--fg);font:16px/1.55 system-ui,-apple-system,sans-serif;padding:24px}
.box{width:100%;max-width:340px;text-align:center}
h1{font-size:26px;margin:0 0 4px;letter-spacing:-.02em}h1 span{color:var(--accent)}
p{color:var(--mut);font-size:14px;margin:0 0 22px}
input{width:100%;padding:15px;font:600 22px/1 system-ui;text-align:center;letter-spacing:.5em;
border:1px solid var(--line);border-radius:10px;background:var(--card);color:var(--fg);margin-bottom:12px}
button{width:100%;padding:14px;border:0;border-radius:10px;background:var(--accent);color:#fff;font:600 16px system-ui}
.err{color:var(--bad);font-size:14px;margin-top:12px}
</style></head><body><div class=box>
<h1>Site<span>Scribe</span></h1>
<p>Structural observation reports from your phone.</p>
<form method=post>
<input name=pin type=tel inputmode=numeric maxlength=12 autocomplete=off
       placeholder="PIN" autofocus aria-label="PIN">
<button>Enter</button>
</form>
{% if error %}<div class=err>{{ error }}</div>{% endif %}
</div></body></html>"""


@app.route("/login", methods=["GET", "POST"])
def login():
    error = None
    ip = _ip()
    if locked(ip):
        error = "Too many attempts. Try again in a few minutes."
    elif request.method == "POST":
        if request.form.get("pin", "").strip() == PIN:
            session.permanent = True
            session["ok"] = True
            _fails.pop(ip, None)
            return redirect(request.args.get("next") or url_for("index"))
        n, _ = _fails.get(ip, (0, 0))
        n += 1
        _fails[ip] = (n, time.time() + LOCK_SECONDS if n >= MAX_FAILS else 0)
        error = "Incorrect PIN."
    return render_template_string(LOGIN_HTML, error=error), (200 if not error else 401)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.after_request
def _no_html_cache(resp):
    """Never let a phone hold a stale page: the UI changes under people mid-visit."""
    if (resp.mimetype or "").startswith("text/html"):
        resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        resp.headers["Pragma"] = "no-cache"
        resp.headers["Expires"] = "0"
    return resp


@app.route("/health")
def health():
    return jsonify(status="ok", service="sitescribe", time=_now())


# ---------------------------------------------------------------- storage

_lock = threading.Lock()


# Stored timestamps stay UTC (unambiguous); anything a person reads is rendered
# in the site's local zone. Override with SITESCRIBE_TZ if the team moves.
LOCAL_TZ_NAME = os.environ.get("SITESCRIBE_TZ", "America/Los_Angeles")

try:
    from zoneinfo import ZoneInfo
    LOCAL_TZ = ZoneInfo(LOCAL_TZ_NAME)
except Exception as exc:  # no tzdata on the box: fall back rather than crash
    print(f"[sitescribe] timezone {LOCAL_TZ_NAME} unavailable ({exc}); using UTC",
          file=sys.stderr)
    LOCAL_TZ = datetime.timezone.utc


def _local(ts=None):
    """A timezone-aware datetime in the site's local zone."""
    dt = (datetime.datetime.now(datetime.timezone.utc) if ts is None
          else datetime.datetime.fromtimestamp(ts, datetime.timezone.utc))
    return dt.astimezone(LOCAL_TZ)


def _now():
    return datetime.datetime.now(datetime.timezone.utc).replace(
        microsecond=0).isoformat().replace("+00:00", "Z")


def _iso(ts):
    return datetime.datetime.fromtimestamp(ts, datetime.timezone.utc).replace(
        microsecond=0).isoformat().replace("+00:00", "Z")


def valid_pid(pid):
    return isinstance(pid, str) and PID_RE.match(pid) is not None


def project_data_dir(pid):
    return os.path.join(DATA_DIR, pid)


def project_uploads(pid):
    return os.path.join(UPLOADS, pid)


def current_pid():
    """The project this browser session has open, or None."""
    pid = session.get("pid")
    if not valid_pid(pid) or not os.path.isdir(project_data_dir(pid)):
        return None
    return pid


def require_project(fn):
    @functools.wraps(fn)
    def wrapper(*a, **kw):
        if current_pid() is None:
            if request.path.startswith("/api/"):
                return jsonify(error="no project open"), 409
            return redirect(url_for("index"))
        return fn(*a, **kw)
    return wrapper


def data_file(pid=None):
    pid = pid or current_pid()
    return os.path.join(project_data_dir(pid), "visit.json")


def plans_dir(pid=None):
    pid = pid or current_pid()
    return os.path.join(project_uploads(pid), "plans")


def item_dir(item_id, pid=None):
    pid = pid or current_pid()
    return os.path.join(project_uploads(pid), secure_filename(item_id))


def _blank():
    return {"items": [], "plans": [], "template": None}


def load(pid=None):
    pid = pid or current_pid()
    path = data_file(pid)
    if not os.path.exists(path):
        return _blank()
    try:
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        if not isinstance(data, dict):
            raise ValueError("visit.json is not an object")
    except (OSError, ValueError) as exc:
        print(f"[sitescribe] {path} unreadable ({exc}); starting empty", file=sys.stderr)
        return _blank()
    data.setdefault("items", [])
    data.setdefault("plans", [])
    data.setdefault("template", None)
    data["items"] = [it for it in data["items"] if isinstance(it, dict)]
    upgraded = needs_upgrade(data)
    if upgraded:
        upgrade_items(data, pid)
    # backfill fields added after some records were written, so an older
    # visit.json never 500s or silently drops data
    for it in data["items"]:
        it.setdefault("item_id", "000")
        it.setdefault("created", _now())
        it["photos"] = [p for p in (it.get("photos") or []) if isinstance(p, dict)]
        for n, p in enumerate(it["photos"], 1):
            p.setdefault("photo_id", f"{it['item_id']}_{n:02d}")
            p["seq"] = photo_seq(p)
            p.setdefault("filename", None)
            p.setdefault("url", None)
            p.setdefault("caption", "")
            p.setdefault("include", True)
            p.setdefault("description", "")
            p.setdefault("importance", "Info")
            p.setdefault("tags", [])
            p.setdefault("location", None)
            p.setdefault("audio_clips", [])
            p.setdefault("created", it["created"])
    if upgraded:
        # persist immediately: the upgrade already moved files on disk, so a
        # read must not leave the json describing the old layout
        try:
            save(data, pid)
        except OSError as exc:
            print(f"[sitescribe] could not persist upgrade for {pid}: {exc}",
                  file=sys.stderr)
    return data


def save(data, pid=None):
    # write to a temp file first so a crash mid-write cannot truncate the visit
    path = data_file(pid)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2)
    os.replace(tmp, path)


def find_item(data, item_id):
    for it in data["items"]:
        if it["item_id"] == item_id:
            return it
    return None


def find_photo(item, photo_id):
    for p in item.get("photos", []):
        if p["photo_id"] == photo_id:
            return p
    return None


def _trailing_num(value):
    """The trailing digits of an id, tolerating old and new formats alike."""
    m = re.search(r"(\d+)\s*$", str(value or ""))
    return int(m.group(1)) if m else 0


def photo_seq(p):
    seq = p.get("seq")
    try:
        return int(seq)
    except (TypeError, ValueError):
        return _trailing_num(str(p.get("photo_id", "")).split("_")[-1])


def next_item_id(data):
    n = max((_trailing_num(it.get("item_id")) for it in data["items"]), default=0) + 1
    if n > MAX_ITEMS:
        raise ApiError(f"this visit already holds {MAX_ITEMS} items; start a new report")
    return f"{n:03d}"


def next_photo_seq(item):
    n = max((photo_seq(p) for p in item.get("photos") or []), default=0) + 1
    if n > MAX_PHOTOS:
        raise ApiError(f"item {item['item_id']} already holds {MAX_PHOTOS} photos; "
                       "start a new item")
    return n


def new_entry(item, *, description="", importance="Info", tags=None, location=None):
    """A fresh photo entry — an observation that may not have its image yet."""
    seq = next_photo_seq(item)
    return {
        "photo_id": f"{item['item_id']}_{seq:02d}",
        "seq": seq,
        "filename": None,
        "url": None,
        "caption": "",
        "include": True,
        "description": description,
        "importance": importance if importance in IMPORTANCE else "Info",
        "tags": list(tags or []),
        "location": dict(location) if location else None,
        "audio_clips": [],
        "created": _now(),
    }


def sorted_entries(item):
    return sorted(item.get("photos") or [], key=photo_seq)


def entry_files(entry):
    """Every file on disk this entry owns: its image and its clips."""
    names = [entry.get("filename")]
    names += [c.get("filename") for c in entry.get("audio_clips") or []]
    return [n for n in names if n]


def next_plan_id(data):
    n = 0
    for p in data["plans"]:
        try:
            n = max(n, int(str(p["plan_id"]).split("-")[-1]))
        except (ValueError, IndexError):
            continue
    return f"PL-{n + 1:02d}"


def sorted_items(data):
    return sorted(data["items"], key=lambda it: it["item_id"])


class ApiError(Exception):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.message = message
        self.status = status


def need_item(data, item_id):
    it = find_item(data, item_id)
    if it is None:
        raise ApiError(f"no such item: {item_id}", 404)
    return it


def need_photo(item, photo_id):
    p = find_photo(item, photo_id)
    if p is None:
        raise ApiError(f"no such photo: {photo_id}", 404)
    return p


# ---------------------------------------------------------------- projects

def load_index():
    if not os.path.exists(PROJECTS_FILE):
        return {"projects": []}
    try:
        with open(PROJECTS_FILE, "r", encoding="utf-8") as fh:
            idx = json.load(fh)
        if not isinstance(idx, dict) or not isinstance(idx.get("projects"), list):
            raise ValueError("projects.json is not an index object")
    except (OSError, ValueError) as exc:
        print(f"[sitescribe] projects.json unreadable ({exc}); starting empty", file=sys.stderr)
        return {"projects": []}
    idx["projects"] = [p for p in idx["projects"]
                       if isinstance(p, dict) and valid_pid(p.get("project_id"))]
    return idx


def save_index(idx):
    tmp = PROJECTS_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(idx, fh, indent=2)
    os.replace(tmp, PROJECTS_FILE)


def find_project(idx, pid):
    for p in idx["projects"]:
        if p.get("project_id") == pid:
            return p
    return None


def need_project(idx, pid):
    # a pid off the wire must never reach a filesystem path unvalidated
    if not valid_pid(pid):
        raise ApiError(f"no such project: {pid}", 404)
    p = find_project(idx, pid)
    if p is None:
        raise ApiError(f"no such project: {pid}", 404)
    return p


def next_pid(idx):
    """Max existing id + 1, so deleting a project never recycles its id."""
    n = 0
    for p in idx["projects"]:
        try:
            n = max(n, int(str(p.get("project_id", "")).split("-")[-1]))
        except (ValueError, IndexError):
            continue
    try:
        for entry in os.listdir(DATA_DIR):
            if PID_RE.match(entry):
                n = max(n, int(entry.split("-")[-1]))
    except OSError:
        pass
    return f"P-{n + 1:03d}"


def _counts(pid):
    """(items, photos, included) for a project; zeros if its file is unusable.

    `included` counts entries now that include lives on the entry, and the
    file is read raw, so an old-shape project reports its item-level flag.
    """
    items = photos = included = 0
    try:
        with open(data_file(pid), "r", encoding="utf-8") as fh:
            data = json.load(fh)
        for it in data["items"]:
            items += 1
            entries = it.get("photos") or []
            photos += len(entries)
            for e in entries:
                if e.get("include", it.get("include", True)):
                    included += 1
    except Exception:
        # one corrupt project must never break the whole listing
        return 0, 0, 0
    return items, photos, included


def project_dict(p):
    pid = p["project_id"]
    created = p.get("created") or _now()
    try:
        updated = _iso(os.path.getmtime(data_file(pid)))
    except OSError:
        updated = created
    items, photos, included = _counts(pid)
    return {"project_id": pid, "name": p.get("name") or pid, "created": created,
            "updated": updated, "items": items, "photos": photos,
            "included": included}


def clean_name(raw, default=None):
    name = ("" if raw is None else str(raw)).strip()
    if not name:
        if default is None:
            raise ApiError("name required")
        return default
    if len(name) > MAX_NAME:
        raise ApiError(f"name must be {MAX_NAME} characters or fewer")
    return name


def auto_name(ts=None):
    return "Site Visit " + _local(ts).strftime("%Y-%m-%d %H:%M")


def make_project(name, created=None, idx=None):
    """Allocate a pid, build its directories and an empty visit, index it."""
    own = idx is None
    if own:
        idx = load_index()
    pid = next_pid(idx)
    os.makedirs(project_data_dir(pid), exist_ok=True)
    os.makedirs(plans_dir(pid), exist_ok=True)
    entry = {"project_id": pid, "name": name, "created": created or _now()}
    save(_blank(), pid)
    idx["projects"].append(entry)
    if own:
        save_index(idx)
    return entry


@app.route("/api/projects")
@require_pin
def api_projects():
    with _lock:
        idx = load_index()
        out = [project_dict(p) for p in idx["projects"]]
    out.sort(key=lambda p: (p["updated"], p["project_id"]), reverse=True)
    return jsonify(projects=out, current=current_pid())


@app.route("/api/projects", methods=["POST"])
@require_pin
def api_project_create():
    body = request.get_json(silent=True) or {}
    name = clean_name(body.get("name"), default=auto_name())
    with _lock:
        entry = make_project(name)
        out = project_dict(entry)
    session["pid"] = entry["project_id"]
    return jsonify(out)


@app.route("/api/projects/close", methods=["POST"])
@require_pin
def api_project_close():
    session.pop("pid", None)
    return jsonify(ok=True)


@app.route("/api/projects/<pid>/open", methods=["POST"])
@require_pin
def api_project_open(pid):
    with _lock:
        entry = need_project(load_index(), pid)
        out = project_dict(entry)
    session["pid"] = pid
    return jsonify(out)


@app.route("/api/projects/<pid>", methods=["PATCH"])
@require_pin
def api_project_patch(pid):
    body = request.get_json(silent=True) or {}
    name = clean_name(body.get("name"))
    with _lock:
        idx = load_index()
        entry = need_project(idx, pid)
        entry["name"] = name
        save_index(idx)
        return jsonify(project_dict(entry))


@app.route("/api/projects/<pid>", methods=["DELETE"])
@require_pin
def api_project_delete(pid):
    with _lock:
        idx = load_index()
        entry = need_project(idx, pid)
        idx["projects"].remove(entry)
        save_index(idx)
        shutil.rmtree(project_data_dir(pid), ignore_errors=True)
        shutil.rmtree(project_uploads(pid), ignore_errors=True)
    if session.get("pid") == pid:
        session.pop("pid", None)
    return jsonify(ok=True)


# ---------------------------------------------------------------- pages

@app.route("/")
@require_pin
def index():
    return render_template("landing.html")


@app.route("/capture")
@require_pin
@require_project
def capture():
    return render_template("index.html")


@app.route("/uploads/<path:relpath>")
@require_pin
def uploaded(relpath):
    # send_from_directory rejects anything that escapes UPLOADS; the project id
    # is already part of relpath, so no project needs to be open to view a file
    return send_from_directory(UPLOADS, relpath)


# ---------------------------------------------------------------- visit / items

@app.route("/api/visit")
@require_pin
@require_project
def api_visit():
    with _lock:
        data = load()
        return jsonify(items=sorted_items(data), plans=data["plans"])


@app.route("/api/item", methods=["POST"])
@require_pin
@require_project
def api_item_create():
    with _lock:
        data = load()
        item = {"item_id": next_item_id(data), "created": _now(), "photos": []}
        # an item is a numbering group; it opens with one entry waiting for
        # an image, so a note can be recorded before the picture is taken
        item["photos"].append(new_entry(item))
        data["items"].append(item)
        os.makedirs(item_dir(item["item_id"]), exist_ok=True)
        save(data)
        return jsonify(item)


def _clean_location(value):
    if value is None:
        return None
    if not isinstance(value, dict):
        raise ApiError("location must be an object or null")
    try:
        x = float(value.get("x"))
        y = float(value.get("y"))
    except (TypeError, ValueError):
        raise ApiError("location x and y must be numbers")
    if not (0.0 <= x <= 1.0 and 0.0 <= y <= 1.0):
        raise ApiError("location x and y must be normalized fractions 0.0-1.0")
    try:
        page = int(value.get("page", 1) or 1)
    except (TypeError, ValueError):
        raise ApiError("location page must be an integer")
    return {"plan_id": str(value.get("plan_id") or ""), "page": page, "x": x, "y": y}


def _clean_tags(value):
    if value is None:
        return []
    if not isinstance(value, list):
        raise ApiError("tags must be a list of strings")
    return [str(t).strip() for t in value if str(t).strip()]


def _clean_importance(value):
    if value not in IMPORTANCE:
        raise ApiError(f"importance must be one of {', '.join(IMPORTANCE)}")
    return value


@app.route("/api/item/<item_id>", methods=["PATCH"])
@require_pin
@require_project
def api_item_patch(item_id):
    # every editable field moved onto the entries; the route stays so the
    # frontend can still read an item back by id
    with _lock:
        data = load()
        return jsonify(need_item(data, item_id))


@app.route("/api/item/<item_id>", methods=["DELETE"])
@require_pin
@require_project
def api_item_delete(item_id):
    with _lock:
        data = load()
        item = need_item(data, item_id)
        data["items"].remove(item)
        shutil.rmtree(item_dir(item_id), ignore_errors=True)
        save(data)
        return jsonify(ok=True)


# ---------------------------------------------------------------- photos

def _inherit_mode():
    """`inherit` off a JSON body, a form field or the query string."""
    body = request.get_json(silent=True)
    raw = (body or {}).get("inherit") if isinstance(body, dict) else None
    if raw is None:
        raw = request.values.get("inherit")
    if raw is None:
        return None
    mode = str(raw).strip().lower()
    if mode not in INHERIT_MODES:
        raise ApiError(f"inherit must be one of {', '.join(INHERIT_MODES)}")
    return mode


def _seed_from_previous(item, mode):
    """Fields a new entry starts with. "new" and "clear" are synonyms."""
    if mode != "keep":
        return {}
    prev = sorted_entries(item)
    if not prev:
        return {}
    last = prev[-1]
    # the observation carries over; the photo and the recordings do not
    return {"description": last.get("description") or "",
            "importance": last.get("importance") or "Info",
            "tags": last.get("tags") or [],
            "location": last.get("location")}


def _attach_photo(entry, f, item_id, pid):
    safe = secure_filename(f.filename) or "photo.jpg"
    filename = f"{entry['photo_id']}_{safe}"
    d = item_dir(item_id, pid)
    os.makedirs(d, exist_ok=True)
    old = entry.get("filename")
    f.save(os.path.join(d, filename))
    if old and old != filename:
        try:
            os.remove(os.path.join(d, old))
        except OSError:
            pass
    entry["filename"] = filename
    entry["url"] = f"/uploads/{pid}/{item_id}/{filename}"


@app.route("/api/item/<item_id>/photo", methods=["POST"])
@require_pin
@require_project
def api_photo_add(item_id):
    f = request.files.get("photo")
    mode = _inherit_mode()
    if (not f or not f.filename) and mode is None:
        raise ApiError("no photo uploaded")
    pid = current_pid()
    # The caller may name the entry the image belongs to. Without it an upload
    # would land on the oldest imageless entry, which is the wrong one as soon
    # as a Keep/Clear entry is waiting ahead of it.
    target_id = (request.form.get("photo_id")
                 or (request.get_json(silent=True) or {}).get("photo_id")
                 or request.args.get("photo_id"))
    with _lock:
        data = load()
        item = need_item(data, item_id)
        entry = None
        if f and f.filename:
            if target_id:
                entry = need_photo(item, target_id)
                if entry.get("filename"):
                    raise ApiError(f"{target_id} already has an image")
            else:
                # newest imageless entry, not the oldest
                entry = next((e for e in reversed(sorted_entries(item))
                              if not e.get("filename")), None)
        if entry is None:
            entry = new_entry(item, **_seed_from_previous(item, mode))
            item.setdefault("photos", []).append(entry)
        if f and f.filename:
            _attach_photo(entry, f, item_id, pid)
        save(data)
        return jsonify(entry)


@app.route("/api/item/<item_id>/photo/<photo_id>", methods=["PATCH"])
@require_pin
@require_project
def api_photo_patch(item_id, photo_id):
    body = request.get_json(silent=True) or {}
    with _lock:
        data = load()
        item = need_item(data, item_id)
        entry = need_photo(item, photo_id)
        if "caption" in body:
            entry["caption"] = "" if body["caption"] is None else str(body["caption"])
        if "description" in body:
            entry["description"] = ("" if body["description"] is None
                                    else str(body["description"]))
        if "importance" in body:
            entry["importance"] = _clean_importance(body["importance"])
        if "tags" in body:
            entry["tags"] = _clean_tags(body["tags"])
        if "include" in body:
            entry["include"] = bool(body["include"])
        if "location" in body:
            entry["location"] = _clean_location(body["location"])
        save(data)
        return jsonify(entry)


@app.route("/api/item/<item_id>/photo/<photo_id>", methods=["DELETE"])
@require_pin
@require_project
def api_photo_delete(item_id, photo_id):
    with _lock:
        data = load()
        item = need_item(data, item_id)
        entry = need_photo(item, photo_id)
        d = item_dir(item_id)
        for name in entry_files(entry):
            try:
                os.remove(os.path.join(d, name))
            except OSError:
                pass  # record removal is what matters; a missing file is not an error
        # an item left with no entries stays: its number is still taken
        item["photos"].remove(entry)
        save(data)
        return jsonify(ok=True)


# ---------------------------------------------------------------- audio

def transcribe(path):
    """Whisper transcription, or "" if unavailable. Never raises."""
    key = os.environ.get("OPENAI_API_KEY")
    if not key:
        print("[sitescribe] no OPENAI_API_KEY; skipping transcription", file=sys.stderr)
        return ""
    try:
        import requests
        with open(path, "rb") as fh:
            r = requests.post(
                "https://api.openai.com/v1/audio/transcriptions",
                headers={"Authorization": f"Bearer {key}"},
                data={"model": "whisper-1"},
                files={"file": (os.path.basename(path), fh, "application/octet-stream")},
                timeout=60,
            )
        if r.status_code != 200:
            print(f"[sitescribe] whisper HTTP {r.status_code}: {r.text[:300]}", file=sys.stderr)
            return ""
        return (r.json().get("text") or "").strip()
    except Exception as exc:
        print(f"[sitescribe] transcription failed: {exc!r}", file=sys.stderr)
        return ""


def next_clip_seq(entry):
    n = 0
    for c in entry.get("audio_clips") or []:
        n = max(n, _trailing_num(c.get("clip_id")))
    return n + 1


@app.route("/api/item/<item_id>/photo/<photo_id>/audio", methods=["POST"])
@require_pin
@require_project
def api_audio(item_id, photo_id):
    f = request.files.get("audio")
    if not f or not f.filename:
        raise ApiError("no audio uploaded")

    # "false"/"0"/"no" means: keep the recording, do not send it anywhere
    want_text = str(request.form.get("transcribe", "true")).lower() not in ("false", "0", "no")
    pid = current_pid()

    with _lock:
        data = load(pid)
        entry = need_photo(need_item(data, item_id), photo_id)
        safe = secure_filename(f.filename) or "audio.webm"
        # keep every recording: <photo_id>_A01_<name>
        clip_id = f"{entry['photo_id']}_A{next_clip_seq(entry):02d}"
        filename = f"{clip_id}_{safe}"
        d = item_dir(item_id, pid)
        os.makedirs(d, exist_ok=True)
        path = os.path.join(d, filename)
        f.save(path)
        size = os.path.getsize(path)

    # transcription can take many seconds — do it outside the lock
    transcript = transcribe(path) if want_text else ""

    with _lock:
        data = load(pid)
        entry = need_photo(need_item(data, item_id), photo_id)
        clip = {
            "clip_id": clip_id,
            "filename": filename,
            "url": f"/uploads/{pid}/{item_id}/{filename}",
            "bytes": size,
            "transcript": transcript,
            "created": _now(),
        }
        entry.setdefault("audio_clips", []).append(clip)
        if transcript and not (entry.get("description") or "").strip():
            entry["description"] = transcript
        save(data, pid)
        return jsonify(clip=clip, transcript=transcript,
                       description=entry.get("description", ""))


@app.route("/api/item/<item_id>/photo/<photo_id>/audio/<clip_id>", methods=["DELETE"])
@require_pin
@require_project
def api_audio_delete(item_id, photo_id, clip_id):
    with _lock:
        data = load()
        entry = need_photo(need_item(data, item_id), photo_id)
        clips = entry.get("audio_clips") or []
        hit = next((c for c in clips if c["clip_id"] == clip_id), None)
        if hit is None:
            raise ApiError(f"no such clip: {clip_id}", 404)
        clips.remove(hit)
        try:
            os.remove(os.path.join(item_dir(item_id), hit["filename"]))
        except OSError:
            pass
        save(data)
        return jsonify(ok=True)


# ---------------------------------------------------------------- plans

def pdf_pages(path):
    try:
        import pypdf
        with open(path, "rb") as fh:
            return max(1, len(pypdf.PdfReader(fh).pages))
    except Exception as exc:
        print(f"[sitescribe] page count unavailable ({exc!r}); assuming 1", file=sys.stderr)
        return 1


@app.route("/api/plan", methods=["POST"])
@require_pin
@require_project
def api_plan():
    f = request.files.get("plan")
    if not f or not f.filename:
        raise ApiError("no plan uploaded")
    pid = current_pid()
    with _lock:
        data = load()
        plan_id = next_plan_id(data)
        safe = secure_filename(f.filename) or "plan.pdf"
        filename = f"{plan_id}_{safe}"
        d = plans_dir(pid)
        os.makedirs(d, exist_ok=True)
        path = os.path.join(d, filename)
        f.save(path)
        # the frontend picks PDF.js vs <img> off `name`, so keep the real extension
        name = os.path.basename(f.filename)
        pages = pdf_pages(path) if name.lower().endswith(".pdf") else 1
        plan = {
            "plan_id": plan_id,
            "name": name,
            "file": filename,
            "url": f"/uploads/{pid}/plans/{filename}",
            "pages": pages,
        }
        data["plans"].append(plan)
        save(data)
        return jsonify(plan)


# ---------------------------------------------------------------- template

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def template_path(pid=None):
    pid = pid or current_pid()
    return os.path.join(project_data_dir(pid), "template.docx")


def has_template(data, pid):
    """A template counts only when both the record and the file are there."""
    return bool(data.get("template")) and os.path.exists(template_path(pid))


def _template_layout(path):
    """How word_report will use the template: "report" (a firm report kept
    as the form, items refilled), "tagged" (Jinja tags) or "styles" (page
    setup and styles only). None when the module cannot say."""
    try:
        from sitescribe import word_report
        return word_report.template_kind(Path(path))
    except Exception as exc:
        print(f"[sitescribe] could not classify template {path}: {exc!r}", file=sys.stderr)
        return None


def _template_json(rec):
    """Flat record for convenience plus `template` so a client can test None."""
    return jsonify(template=rec, **(rec or {}))


@app.route("/api/template")
@require_pin
@require_project
def api_template_get():
    with _lock:
        data = load()
        return _template_json(data.get("template"))


@app.route("/api/template", methods=["POST"])
@require_pin
@require_project
def api_template_upload():
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
    except ImportError:
        return jsonify(error="python-docx not installed"), 501

    f = request.files.get("template")
    if not f or not f.filename:
        raise ApiError("no template uploaded")
    name = os.path.basename(f.filename)
    low = name.lower()
    if low.endswith(".doc"):
        raise ApiError("legacy .doc files are not supported — open it in Word "
                       "and use Save As to make a .docx, then upload that")
    if not low.endswith(".docx"):
        raise ApiError("template must be a Word .docx file")

    pid = current_pid()
    with _lock:
        path = template_path(pid)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        # land it beside the real name first: a file that will not open must
        # not replace a template that already works
        tmp = path + ".tmp"
        f.save(tmp)
        try:
            doc = Document(tmp)
            styles = [s.name for s in doc.styles if s.type == WD_STYLE_TYPE.PARAGRAPH]
        except Exception as exc:
            print(f"[sitescribe] rejected template {name!r}: {exc!r}", file=sys.stderr)
            try:
                os.remove(tmp)
            except OSError:
                pass
            raise ApiError("not a readable .docx file")
        os.replace(tmp, path)
        data = load(pid)
        rec = {"name": name, "uploaded": _now(), "styles": styles,
               "layout": _template_layout(path)}
        data["template"] = rec
        save(data, pid)
        return _template_json(rec)


@app.route("/api/template", methods=["DELETE"])
@require_pin
@require_project
def api_template_delete():
    pid = current_pid()
    with _lock:
        try:
            os.remove(template_path(pid))
        except OSError:
            pass  # record removal is what matters; a missing file is not an error
        data = load(pid)
        data["template"] = None
        save(data, pid)
        return jsonify(ok=True)


# ---------------------------------------------------------------- export

def _sibling_module(name):
    """webapp/<name>.py, loaded from beside this file and cached in
    sys.modules. Going by path rather than by name works however the server
    was started: run.py, `python server.py`, gunicorn, or a wsgi file that
    loads server.py by path (where neither `webapp` nor webapp/ is on
    sys.path)."""
    key = f"sitescribe_webapp_{name}"
    mod = sys.modules.get(key)
    if mod is None:
        import importlib.util
        spec = importlib.util.spec_from_file_location(key, os.path.join(BASE, f"{name}.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        sys.modules[key] = mod
    return mod


def _visit_date(entry):
    """The project's creation time in the site's zone, else now."""
    try:
        when = datetime.datetime.fromisoformat(str(entry.get("created")).replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return _local()
    return when.astimezone(LOCAL_TZ) if when.tzinfo else when


def _build_project(data, pid, entry, work_dir):
    """The visit as a models.Project via bridge.to_project(), plan pages
    rendered into work_dir (which must outlive every use of the Project).
    (project, None), or (None, reason) when the adapter cannot run."""
    try:
        bridge = _sibling_module("bridge")
    except (ImportError, OSError) as exc:
        return None, _fallback(f"the visit adapter could not load ({exc})")
    try:
        return bridge.to_project(
            data,
            name=entry.get("name") or pid,
            item_dir=lambda item_id: item_dir(item_id, pid),
            plans_dir=plans_dir(pid),
            work_dir=Path(work_dir) / "plans",
            template=template_path(pid) if has_template(data, pid) else None,
            visit_date=_visit_date(entry),
            tz=LOCAL_TZ,
        ), None
    except Exception as exc:
        return None, _fallback(f"the visit could not be converted ({type(exc).__name__}: {exc})")


def _module_report(project, work_dir):
    """The report from sitescribe.word_report. (buffer, None) on success;
    (None, reason) when the module is not built yet (NotImplementedError) or
    cannot run, so the built-in report below takes over and the user still
    gets a file. The reason is short enough to show the user; the traceback
    goes to stderr. Raises ImportError only when python-docx itself is
    missing."""
    try:
        from sitescribe import word_report
    except ImportError as exc:
        return None, _fallback(f"the Word report module could not load ({exc})")
    try:
        out = word_report.build_report(project, Path(work_dir) / "site_visit_report.docx")
        return io.BytesIO(Path(out).read_bytes()), None
    except NotImplementedError:
        return None, _fallback("the Word report module is not written yet", trace=False)
    except ImportError as exc:
        if "python-docx" in str(exc):
            raise
        return None, _fallback(f"the Word report module needs a package ({exc})", trace=False)
    except Exception as exc:
        return None, _fallback(f"the Word report module failed ({type(exc).__name__}: {exc})")


def _fallback(reason, *, trace=True):
    """Log why the built-in report is shipping instead of the module's, with
    the traceback when there is one, and return the reason for the user."""
    reason = " ".join(str(reason).split())
    print(f"[sitescribe] {reason}; exporting built-in report", file=sys.stderr)
    if trace:
        traceback.print_exc(file=sys.stderr)
    return reason


def styled(doc, text, *candidates):
    """Add a paragraph using the first style the document actually has.

    Uploaded templates differ wildly in which named styles they define, so a
    missing one steps to the next candidate and finally to Normal. The style is
    applied after the paragraph exists — add_paragraph(style=...) inserts the
    text before it validates the name, which on failure would strand a
    duplicate. ValueError covers a template where the name exists but belongs
    to a character or table style.
    """
    p = doc.add_paragraph(text)
    for name in candidates:
        try:
            p.style = name
            break
        except (KeyError, ValueError):
            continue
    return p


def _strip_body(doc):
    """Empty the template's body but keep its trailing <w:sectPr>.

    sectPr holds page size, margins and the header/footer references, so
    dropping it would lose the letterhead. Everything before it is the
    template's own prose — a user may well upload a filled-in past report
    rather than a blank shell, and that content must not reach the output.
    """
    body = doc.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)


def _write_report(doc, data, pid, *, title, heading, body, caption, Inches):
    """The report itself: same content and ordering for either export path.

    One section per included entry, in item then seq order — the entry is the
    observation, and an entry with no image still has something to report.
    """
    plan_names = {p["plan_id"]: p.get("name", p["plan_id"]) for p in data["plans"]}

    title("Site Visit Report")
    # "%-I" (hour without a leading zero) is glibc-only; Windows strftime rejects it
    now = _local()
    body(f"{now:%B %d, %Y} at {now.hour % 12 or 12}:{now:%M %p %Z}")

    for it in sorted_items(data):
        for e in sorted_entries(it):
            if not e.get("include", True):
                continue
            heading(f"{e['photo_id']} — {e.get('importance', 'Info')}")
            if (e.get("description") or "").strip():
                body(e["description"])
            tags = e.get("tags") or []
            if tags:
                body("Tags: " + ", ".join(tags))
            loc = e.get("location")
            if loc:
                name = plan_names.get(loc.get("plan_id"), loc.get("plan_id") or "unknown")
                body(f"Plan {name}, page {loc.get('page', 1)} "
                     f"(x={float(loc.get('x', 0)):.2f}, y={float(loc.get('y', 0)):.2f})")
            if not e.get("filename"):
                continue
            path = os.path.join(item_dir(it["item_id"], pid), e["filename"])
            if not os.path.exists(path):
                continue
            try:
                doc.add_picture(path, width=Inches(4.5))
            except Exception as exc:
                print(f"[sitescribe] could not embed {path}: {exc!r}", file=sys.stderr)
                continue
            if (e.get("caption") or "").strip():
                caption(e["caption"])


def _italic_caption(doc, text, Pt):
    cap = doc.add_paragraph()
    run = cap.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    return cap


def _styled_caption(doc, text, Pt):
    """Caption style when the template has one, hand-italicised otherwise."""
    cap = doc.add_paragraph(text)
    try:
        cap.style = "Caption"
    except (KeyError, ValueError):
        for run in cap.runs:
            run.italic = True
            run.font.size = Pt(9)
    return cap


def _plain_doc(data, pid, Document, Inches, Pt):
    """The original report, built on python-docx's own default template."""
    doc = Document()
    _write_report(
        doc, data, pid,
        title=lambda t: doc.add_heading(t, level=0),
        heading=lambda t: doc.add_heading(t, level=1),
        body=doc.add_paragraph,
        caption=lambda t: _italic_caption(doc, t, Pt),
        Inches=Inches,
    )
    return doc


def _styled_doc(data, pid, path, Document, Inches, Pt):
    """The same report written into the firm's template: styles, margins,
    headers and footers all come along because the template *is* the base."""
    doc = Document(path)
    _strip_body(doc)
    _write_report(
        doc, data, pid,
        title=lambda t: styled(doc, t, "Title", "Heading 1"),
        heading=lambda t: styled(doc, t, "Heading 1", "Heading 2"),
        body=lambda t: styled(doc, t, "Body Text", "Normal"),
        caption=lambda t: _styled_caption(doc, t, Pt),
        Inches=Inches,
    )
    return doc


def _to_buffer(doc):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _report_and_project(data, pid, entry, work_dir):
    """(buffer, notice, project): the Word report, sitescribe.word_report's
    when it can run, else the built-in one below, in the template if there is
    one, else plain. notice is None when the module's report shipped,
    otherwise one line saying why the built-in one did, for the user to see.
    project is the Project the report was built from (None when the visit
    could not be converted), for the bundle to reuse: the plan pages it
    carries are rendered into work_dir and live as long as it does.
    Raises ImportError when python-docx is missing."""
    if entry is None:
        entry = find_project(load_index(), pid) or {}
    project, notice = _build_project(data, pid, entry, work_dir)
    if project is not None:
        buf, notice = _module_report(project, work_dir)
        if buf is not None:
            return buf, None, project
    notice = f"built-in report layout used because {notice}"

    from docx import Document
    from docx.shared import Inches, Pt

    if has_template(data, pid):
        path = template_path(pid)
        try:
            return _to_buffer(_styled_doc(data, pid, path, Document, Inches, Pt)), notice, project
        except Exception as exc:
            # a template gone bad must never stand between the user and a
            # report — say so on the log and fall back to the plain one
            print(f"[sitescribe] template {path} unusable ({exc!r}); "
                  "exporting plain report", file=sys.stderr)
            notice += f"; the template could not be used ({type(exc).__name__})"
    return _to_buffer(_plain_doc(data, pid, Document, Inches, Pt)), notice, project


def _report_buffer(data, pid, entry=None):
    """(buffer, notice) for the report-only download."""
    with tempfile.TemporaryDirectory() as tmp:
        buf, notice, _ = _report_and_project(data, pid, entry, tmp)
    return buf, notice


def _notice_headers(notice):
    """The fallback notice as a response header the page can show. Header
    values are latin-1 and one line, so anything else becomes '?'."""
    if not notice:
        return {}
    clean = "".join(ch if " " <= ch <= "~" else "?" for ch in notice)[:300]
    return {"X-SiteScribe-Report-Notice": clean}


def _export_snapshot(pid):
    """The visit and its index entry, read under the lock. The report itself
    is built outside it: rendering plans and photos takes seconds on a big
    visit and must not stall every other request meanwhile."""
    with _lock:
        return load(pid), (find_project(load_index(), pid) or {})


@app.route("/api/export/docx")
@require_pin
@require_project
def api_export_docx():
    pid = current_pid()
    data, entry = _export_snapshot(pid)
    try:
        buf, notice = _report_buffer(data, pid, entry)
    except ImportError:
        return jsonify(error="python-docx not installed"), 501

    resp = send_file(
        buf,
        mimetype=DOCX_MIME,
        as_attachment=True,
        download_name="site_visit_report.docx",
    )
    resp.headers.update(_notice_headers(notice))
    return resp


def _export_bundle():
    """webapp/export_bundle.py, loaded from beside this file (see
    _sibling_module)."""
    return _sibling_module("export_bundle")


@app.route("/api/export/bundle")
@require_pin
@require_project
def api_export_bundle():
    """A zip of one folder: the Word report, photos/ with each entry written
    into the photo's EXIF, and plans/ with every pinned plan page. Built as
    the project's latest export (the same file /api/export/build makes) and
    served from disk, so the zip is never held in memory."""
    pid = current_pid()
    try:
        path, name, notice = _build_export(pid)
    except ImportError as exc:
        return jsonify(error=str(exc)), 501
    resp = send_file(path, mimetype="application/zip", as_attachment=True,
                     download_name=name, conditional=False, max_age=0)
    resp.headers.update(_notice_headers(notice))
    return resp


def _write_bundle(pid, out_path):
    """Write the project's bundle zip to out_path; (file name, fallback
    notice or None). The visit is converted and its plan pages rendered once,
    for the report and the bundle both, and the zip goes straight to disk.
    Raises ImportError with a message for the user when a package is missing."""
    data, entry = _export_snapshot(pid)
    name = entry.get("name") or pid
    try:
        bundle = _export_bundle()
    except (ImportError, OSError) as exc:
        print(f"[sitescribe] export bundle unavailable: {exc!r}", file=sys.stderr)
        raise ImportError(f"photo/plan export unavailable: {exc}") from None
    with tempfile.TemporaryDirectory() as tmp:
        try:
            report, notice, project = _report_and_project(data, pid, entry, tmp)
        except ImportError:
            raise ImportError("python-docx not installed") from None
        if project is None:
            # the zip can only hold what the Project carries: say so, rather
            # than hand over a bundle that silently lacks the photos and plans
            notice = f"{notice}; the photos and plan pages were left out of the zip"
        bundle.write_zip(project, report.getvalue(), name=name,
                         out_path=Path(out_path), work_dir=Path(tmp) / "bundle")
    return f"{bundle.folder_name(name)}.zip", notice


def _build_export(pid):
    """Build the project's bundle as its latest export (export.zip beside the
    visit, with a .json naming it); (path, file name, notice). Written to a
    .tmp beside it and renamed in, so a download never sees a half-built
    file. Raises ImportError with a message for the user when a package is
    missing."""
    path = export_file(pid)
    tmp = path + ".tmp"
    os.makedirs(os.path.dirname(path), exist_ok=True)
    try:
        name, notice = _write_bundle(pid, tmp)
    except BaseException:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise
    os.replace(tmp, path)
    with open(path + ".json", "w", encoding="utf-8") as fh:
        json.dump({"name": name, "built": _now()}, fh)
    return path, name, notice


def export_file(pid=None):
    pid = pid or current_pid()
    return os.path.join(project_data_dir(pid), "export.zip")


@app.route("/api/export/build", methods=["POST"])
@require_pin
@require_project
def api_export_build():
    """Build the bundle and keep it as the project's latest export, so the
    page can show progress and errors (fetch) and then download a file from
    disk by plain navigation, which works the same in every phone browser.
    Returns where to download it and the fallback notice, if any."""
    pid = current_pid()
    started = time.time()
    try:
        path, name, notice = _build_export(pid)
    except ImportError as exc:
        return jsonify(error=str(exc)), 501
    return jsonify(url="/api/export/download", name=name, notice=notice,
                   bytes=os.path.getsize(path), seconds=round(time.time() - started, 1))


@app.route("/api/export/download")
@require_pin
@require_project
def api_export_download():
    """The latest bundle built by /api/export/build."""
    pid = current_pid()
    path = export_file(pid)
    if not os.path.exists(path):
        raise ApiError("no export built yet", 404)
    name = "site_visit_report.zip"
    try:
        with open(path + ".json", encoding="utf-8") as fh:
            name = json.load(fh).get("name") or name
    except (OSError, ValueError):
        pass
    return send_file(path, mimetype="application/zip", as_attachment=True,
                     download_name=name, conditional=False, max_age=0)


# ---------------------------------------------------------------- errors

@app.errorhandler(ApiError)
def handle_api_error(exc):
    return jsonify(error=exc.message), exc.status


@app.errorhandler(Exception)
def handle_any(exc):
    if isinstance(exc, HTTPException):
        if request.path.startswith("/api/"):
            return jsonify(error=exc.description), exc.code
        return exc
    print(f"[sitescribe] unhandled error on {request.path}: {exc!r}", file=sys.stderr)
    if request.path.startswith("/api/"):
        return jsonify(error="internal server error"), 500
    return "Internal Server Error", 500


# ---------------------------------------------------------------- migration

# the old ids: item "I-003", photo "I-003-P-02", audio file "I-003_audio_01_x"
_OLD_PREFIX_RE = re.compile(r"^I-\d+(?:-P-\d+)?(?:_audio_\d+)?_")


def needs_upgrade(data):
    """True while any item still carries the pre-entry shape.

    Detected by an `I-` item id, or by a photo record missing `description`
    — the field that moved down from the item onto the entry.
    """
    for it in data.get("items") or []:
        if "I-" in str(it.get("item_id", "")):
            return True
        photos = [p for p in (it.get("photos") or []) if isinstance(p, dict)]
        if photos and "description" not in photos[0]:
            return True
    return False


def _upgrade_item_dir(pid, old_id, new_id):
    """Move uploads/<pid>/I-003 to uploads/<pid>/003.

    Returns (directory name that actually holds the files, its path or None).
    """
    if not pid:
        return old_id, None
    old_dir = os.path.join(project_uploads(pid), secure_filename(old_id))
    new_dir = os.path.join(project_uploads(pid), secure_filename(new_id))
    if old_id == new_id or not os.path.isdir(old_dir):
        return new_id, new_dir
    try:
        if os.path.isdir(new_dir):  # a half-finished earlier attempt
            for entry in os.listdir(old_dir):
                shutil.move(os.path.join(old_dir, entry), os.path.join(new_dir, entry))
            os.rmdir(old_dir)
        else:
            os.replace(old_dir, new_dir)
    except OSError as exc:
        print(f"[sitescribe] could not move {old_dir} -> {new_dir} ({exc!r}); "
              "leaving the files where they are", file=sys.stderr)
        return old_id, old_dir
    return new_id, new_dir


def _reprefix(d, name, prefix):
    """Re-prefix one stored upload. Keeps the old name if the rename fails —
    a record pointing at a real file beats a record with a tidy id."""
    if not name or d is None or name.startswith(prefix + "_"):
        return name
    tail = _OLD_PREFIX_RE.sub("", name)
    new = f"{prefix}_{tail}" if tail else prefix
    src, dst = os.path.join(d, name), os.path.join(d, new)
    if not os.path.exists(src):
        # this runs on every load until a save persists the new shape, so the
        # rename may already be done; otherwise the file is simply gone
        return new if os.path.exists(dst) else name
    try:
        os.replace(src, dst)
    except OSError as exc:
        print(f"[sitescribe] could not rename {name} to {new} in {d} ({exc!r}); "
              "keeping the old filename", file=sys.stderr)
        return name
    return new


def upgrade_items(data, pid):
    """Move the observation down from each item onto its photo entries.

    The item's one description described all of its photos, so every entry
    inherits it; the ids lose their `I-`/`-P-` prefixes and the files on disk
    are re-prefixed to match. Memory only — the next save() persists it.
    """
    for it in data["items"]:
        old_id = str(it.get("item_id", ""))
        new_id = f"{_trailing_num(old_id):03d}"
        it["item_id"] = new_id
        created = it.get("created") or _now()
        it["created"] = created
        dir_name, d = _upgrade_item_dir(pid, old_id, new_id)

        shared = {
            "include": bool(it.pop("include", True)),
            "description": it.pop("description", "") or "",
            "importance": it.pop("importance", "Info") or "Info",
            "tags": [str(t) for t in (it.pop("tags", None) or [])],
            "location": it.pop("location", None),
        }
        clips = [c for c in (it.pop("audio_clips", None) or []) if isinstance(c, dict)]
        it.pop("audio", None)  # the legacy single-clip pointer is gone

        entries, used = [], set()
        for n, p in enumerate([q for q in (it.get("photos") or [])
                               if isinstance(q, dict)], 1):
            seq = _trailing_num(str(p.get("photo_id", "")).split("-")[-1]) or n
            if seq in used:
                seq = max(used, default=0) + 1
            used.add(seq)
            photo_id = f"{new_id}_{seq:02d}"
            fname = _reprefix(d, p.get("filename"), photo_id)
            entries.append(_upgraded_entry(
                photo_id, seq, shared, pid, dir_name, fname,
                caption=p.get("caption") or "", url=p.get("url"),
                created=p.get("created") or created))

        if not entries:  # an item with no photos still held an observation
            entries.append(_upgraded_entry(f"{new_id}_01", 1, shared, pid,
                                           dir_name, None, created=created))

        # A clip file can only carry one name. Giving every entry a record that
        # points at the same file leaves dangling records as soon as one is
        # deleted, so the clips attach to the first entry only.
        first = entries[0]
        first["audio_clips"] = [
            {"clip_id": f"{first['photo_id']}_A{n:02d}",
             "filename": _reprefix(d, c.get("filename"),
                                   f"{first['photo_id']}_A{n:02d}"),
             "url": None,
             "bytes": int(c.get("bytes") or 0),
             "transcript": c.get("transcript") or "",
             "created": c.get("created") or first["created"]}
            for n, c in enumerate(clips, 1)]
        for clip in first["audio_clips"]:
            clip["url"] = (f"/uploads/{pid}/{dir_name}/{clip['filename']}"
                           if pid and clip["filename"] else None)
        for e in entries[1:]:
            e["audio_clips"] = []

        it["photos"] = entries
    return data


def _upgraded_entry(photo_id, seq, shared, pid, dir_name, fname, *, caption="",
                    url=None, created=None):
    loc = shared["location"]
    return {
        "photo_id": photo_id,
        "seq": seq,
        "filename": fname,
        "url": (f"/uploads/{pid}/{dir_name}/{fname}" if (pid and fname)
                else (url if fname else None)),
        "caption": caption,
        "include": shared["include"],
        "description": shared["description"],
        "importance": shared["importance"],
        "tags": list(shared["tags"]),
        "location": dict(loc) if isinstance(loc, dict) else None,
        "audio_clips": [],
        "created": created or _now(),
    }


def migrate_legacy():
    """Fold a pre-projects single visit into its own project. Idempotent."""
    if not os.path.exists(LEGACY_DATA_FILE):
        return None
    with _lock:
        if not os.path.exists(LEGACY_DATA_FILE):
            return None
        mtime = os.path.getmtime(LEGACY_DATA_FILE)
        idx = load_index()
        pid = next_pid(idx)
        os.makedirs(project_data_dir(pid), exist_ok=True)
        os.makedirs(project_uploads(pid), exist_ok=True)
        dest = os.path.join(project_data_dir(pid), "visit.json")
        os.replace(LEGACY_DATA_FILE, dest)

        moved = []
        for entry in sorted(os.listdir(UPLOADS)):
            src = os.path.join(UPLOADS, entry)
            # leave already-migrated project trees alone
            if not os.path.isdir(src) or PID_RE.match(entry):
                continue
            shutil.move(src, os.path.join(project_uploads(pid), entry))
            moved.append(entry)
        os.makedirs(plans_dir(pid), exist_ok=True)

        _reprefix_urls(dest, pid, mtime)

        idx["projects"].append({"project_id": pid, "name": auto_name(mtime),
                                "created": _iso(mtime)})
        save_index(idx)
    print(f"[sitescribe] migrated legacy visit.json into {pid} "
          f"({auto_name(mtime)}); moved {len(moved)} upload dir(s): "
          f"{', '.join(moved) or 'none'}", file=sys.stderr)
    return pid


def _reprefix_urls(path, pid, mtime):
    """Legacy /uploads/<item>/... urls now live one level deeper."""
    try:
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
    except (OSError, ValueError) as exc:
        print(f"[sitescribe] could not rewrite urls in {path}: {exc}", file=sys.stderr)
        return

    def fix(rec):
        url = rec.get("url")
        if isinstance(url, str) and url.startswith("/uploads/") and \
                not url.startswith(f"/uploads/{pid}/"):
            rec["url"] = "/uploads/" + pid + url[len("/uploads"):]

    for it in data.get("items") or []:
        if not isinstance(it, dict):
            continue
        for rec in (it.get("photos") or []) + (it.get("audio_clips") or []):
            if isinstance(rec, dict):
                fix(rec)
    for pl in data.get("plans") or []:
        if isinstance(pl, dict):
            fix(pl)

    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2)
    os.replace(tmp, path)
    os.utime(path, (mtime, mtime))  # keep "updated" honest for the listing


migrate_legacy()


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
